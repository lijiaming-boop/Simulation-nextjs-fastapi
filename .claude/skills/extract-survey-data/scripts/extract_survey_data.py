#!/usr/bin/env python3
"""
从中央空调节能调研表 Excel 中提取月度运行数据
"""
import sys
import json
import re
from pathlib import Path
from typing import Dict, List, Optional, Any

try:
    import openpyxl
except ImportError:
    print("错误: 需要安装 openpyxl 库")
    print("请运行: pip install openpyxl")
    sys.exit(1)


def parse_frequency(text: Any) -> Optional[float]:
    """解析频率字段，如 '工频50HZ', '45Hz', '50' 等"""
    if text is None or text == "":
        return None

    text_str = str(text).upper()
    # 提取数字部分
    match = re.search(r'(\d+\.?\d*)', text_str)
    if match:
        try:
            return float(match.group(1))
        except ValueError:
            return None
    return None


def parse_temperature_diff(text: Any) -> Optional[float]:
    """解析温差字段，如 '3-5℃', '4℃', '4.5' 等"""
    if text is None or text == "":
        return None

    text_str = str(text)
    # 移除 ℃ 符号
    text_str = text_str.replace('℃', '').replace('°C', '').strip()

    # 处理区间格式 '3-5' 或 '3~5'
    if '-' in text_str or '~' in text_str:
        parts = re.split(r'[-~]', text_str)
        try:
            nums = [float(p.strip()) for p in parts if p.strip()]
            if len(nums) == 2:
                return (nums[0] + nums[1]) / 2  # 返回中间值
        except ValueError:
            pass

    # 处理 '4℃/5℃' 格式，取第一个值
    if '/' in text_str:
        parts = text_str.split('/')
        try:
            return float(parts[0].strip())
        except ValueError:
            pass

    # 尝试直接解析为数字
    try:
        return float(text_str)
    except ValueError:
        return None


def parse_temperature(text: Any) -> Optional[float]:
    """解析温度字段"""
    if text is None or text == "":
        return None

    text_str = str(text).replace('℃', '').replace('°C', '').strip()
    try:
        return float(text_str)
    except ValueError:
        return None


def parse_energy(text: Any) -> Optional[float]:
    """解析能耗字段"""
    if text is None or text == "":
        return None

    try:
        return float(text)
    except (ValueError, TypeError):
        return None


def validate_data(monthly_data: List[Dict]) -> bool:
    """验证数据质量"""
    # 检查至少有 3 个月有能耗数据
    valid_months = sum(1 for m in monthly_data if m.get('total_energy') is not None)
    if valid_months < 3:
        return False

    # 验证温度和频率范围
    for month_data in monthly_data:
        # 检查温度范围
        for temp_field in ['chiller_supply_temp', 'cooling_tower_outlet_temp']:
            temp = month_data.get(temp_field)
            if temp is not None and not (-20 <= temp <= 50):
                return False

        # 检查频率范围
        for freq_field in ['cooling_pump_freq', 'freezing_pump_freq']:
            freq = month_data.get(freq_field)
            if freq is not None and not (0 <= freq <= 50):
                return False

    return True


def find_data_sheet(workbook) -> Optional[Any]:
    """智能识别包含月度数据的工作表"""
    # 优先查找名称包含"冷站"的工作表
    for sheet_name in workbook.sheetnames:
        if '冷站' in sheet_name:
            return workbook[sheet_name]

    # 如果没找到，返回第一个工作表
    if len(workbook.sheetnames) > 0:
        return workbook[workbook.sheetnames[0]]

    return None


def extract_region_from_address(sheet) -> Optional[str]:
    """从工作表中尝试提取项目地区"""
    regions = ['华南', '华东', '华中', '西南', '西北', '华北']

    # 搜索前 20 行中是否包含地区信息
    for row in sheet.iter_rows(min_row=1, max_row=20, values_only=True):
        for cell in row:
            if cell and isinstance(cell, str):
                for region in regions:
                    if region in cell:
                        return region

    return None


def extract_monthly_data(sheet, row_mappings: Dict[str, int]) -> List[Dict]:
    """提取 12 个月的数据"""
    monthly_data = []

    # 假设月度数据从列 B (索引 2) 开始，每月一列
    # 需要灵活识别列位置
    for month in range(1, 13):
        col_idx = month + 1  # 假设第一列是标题，月份数据从第 2 列开始

        month_record = {
            'month': month,
            'total_energy': parse_energy(sheet.cell(row=row_mappings.get('total_energy', 47), column=col_idx).value),
            'chiller_supply_temp': parse_temperature(sheet.cell(row=row_mappings.get('chiller_supply_temp', 38), column=col_idx).value),
            'cooling_pump_freq': parse_frequency(sheet.cell(row=row_mappings.get('cooling_pump_freq', 39), column=col_idx).value),
            'cooling_water_temp_diff': parse_temperature_diff(sheet.cell(row=row_mappings.get('cooling_water_temp_diff', 40), column=col_idx).value),
            'freezing_pump_freq': parse_frequency(sheet.cell(row=row_mappings.get('freezing_pump_freq', 41), column=col_idx).value),
            'freezing_water_temp_diff': parse_temperature_diff(sheet.cell(row=row_mappings.get('freezing_water_temp_diff', 42), column=col_idx).value),
            'cooling_tower_outlet_temp': parse_temperature(sheet.cell(row=row_mappings.get('cooling_tower_outlet_temp', 44), column=col_idx).value),
        }

        monthly_data.append(month_record)

    return monthly_data


def extract_survey_data(excel_path: str, output_path: str) -> Dict:
    """主提取函数"""
    try:
        # 读取 Excel 文件
        workbook = openpyxl.load_workbook(excel_path, data_only=True)

        # 查找数据工作表
        sheet = find_data_sheet(workbook)
        if sheet is None:
            return {
                'error': True,
                'message': '未能识别调研表中的工作表，请确认上传了正确的文件'
            }

        # 提取项目地区
        project_region = extract_region_from_address(sheet)

        # 行号映射（默认值，可根据实际情况调整）
        row_mappings = {
            'chiller_supply_temp': 38,
            'cooling_pump_freq': 39,
            'cooling_water_temp_diff': 40,
            'freezing_pump_freq': 41,
            'freezing_water_temp_diff': 42,
            'cooling_tower_outlet_temp': 44,
            'total_energy': 47,
        }

        # 提取月度数据
        monthly_data = extract_monthly_data(sheet, row_mappings)

        # 验证数据质量
        if not validate_data(monthly_data):
            return {
                'error': True,
                'message': '未能识别调研表中的月度运行数据，请确认上传了正确的文件'
            }

        # 构建结果
        result = {
            'project_region': project_region,
            'monthly_data': monthly_data
        }

        # 输出 JSON
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(result, f, ensure_ascii=False, indent=2)

        return result

    except Exception as e:
        return {
            'error': True,
            'message': f'Excel 文件无法读取: {str(e)}'
        }


def main():
    if len(sys.argv) != 3:
        print("用法: python extract_survey_data.py <survey_excel_path> <output_json_path>")
        sys.exit(1)

    excel_path = sys.argv[1]
    output_path = sys.argv[2]

    # 检查输入文件是否存在
    if not Path(excel_path).exists():
        print(f"错误: 文件不存在 - {excel_path}")
        sys.exit(1)

    # 执行提取
    result = extract_survey_data(excel_path, output_path)

    # 如果有错误，打印并退出
    if result.get('error'):
        print(f"提取失败: {result.get('message')}")
        sys.exit(1)

    print(f"✓ 数据提取成功，输出到: {output_path}")
    print(f"  - 项目地区: {result.get('project_region', '未识别')}")

    valid_months = sum(1 for m in result['monthly_data'] if m.get('total_energy') is not None)
    print(f"  - 有效月份数: {valid_months}/12")


if __name__ == '__main__':
    main()
