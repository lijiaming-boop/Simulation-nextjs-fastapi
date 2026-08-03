#!/usr/bin/env python3
"""
空压机调研表文本提取器
从Excel、Word、PPT等文件中提取文本内容
"""

import json
import sys
import os
import re
from pathlib import Path
from typing import Dict, List, Any, Optional

# 添加项目根目录到Python路径
project_root = Path(__file__).parent.parent.parent.parent
sys.path.insert(0, str(project_root))

try:
    import pandas as pd
except ImportError:
    print("错误: 需要安装 pandas 库")
    print("请运行: pip install pandas")
    sys.exit(1)

try:
    import openpyxl
except ImportError:
    print("错误: 需要安装 openpyxl 库")
    print("请运行: pip install openpyxl")
    sys.exit(1)

try:
    from docx import Document
    from docx.table import Table
except ImportError:
    print("警告: python-docx 未安装，Word文档读取功能将不可用")
    Document = None

try:
    from pptx import Presentation
except ImportError:
    print("警告: python-pptx 未安装，PowerPoint文档读取功能将不可用")
    Presentation = None

try:
    import textract
    HAS_TEXTRACT = True
except ImportError:
    print("警告: textract 未安装，老式Word(.doc)文档读取功能将不可用")
    print("请运行: pip install textract")
    print("同时需要安装 antiword: https://www.softpedia.com/get/Office-tools/Other-Office-Tools/Antiword.shtml")
    HAS_TEXTRACT = False


class TextExtractor:
    """文本提取器 - 从各种文件格式中提取文本内容"""
    
    def __init__(self):
        """初始化文本提取器"""
        self.extracted_text = ""
        self.file_type = ""
    
    def extract_from_file(self, file_path: str) -> str:
        """从文件中提取文本内容"""
        
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")
        
        # 确定文件类型
        file_ext = file_path.suffix.lower()
        self.file_type = file_ext
        
        print(f"📄 [文本提取] 开始提取文件: {file_path.name}")
        print(f"📄 [文本提取] 文件类型: {file_ext}")
        
        # 根据文件类型选择提取方法
        if file_ext in ['.xlsx', '.xls']:
            text = self._extract_from_excel(file_path)
        elif file_ext in ['.docx', '.doc']:
            text = self._extract_from_word(file_path)
        elif file_ext in ['.pptx', '.ppt']:
            text = self._extract_from_ppt(file_path)
        else:
            raise ValueError(f"不支持的文件格式: {file_ext}")
        
        self.extracted_text = text
        print(f"✅ [文本提取] 提取完成，文本长度: {len(text)} 字符")
        
        return text
    
    def _extract_from_excel(self, file_path: Path) -> str:
        """从Excel文件中提取文本"""
        
        text_parts = []
        
        try:
            # 尝试使用openpyxl引擎
            df = pd.read_excel(file_path, engine='openpyxl', sheet_name=None)
            
            for sheet_name, sheet_data in df.items():
                text_parts.append(f"=== 工作表: {sheet_name} ===")
                
                # 提取表头
                if not sheet_data.empty:
                    headers = [str(col) for col in sheet_data.columns]
                    text_parts.append("表头: " + " | ".join(headers))
                    
                    # 提取数据行
                    for _, row in sheet_data.iterrows():
                        row_text = " | ".join([str(val) if pd.notna(val) else "" for val in row])
                        text_parts.append(row_text)
                
                text_parts.append("")  # 空行分隔
        
        except Exception as e:
            print(f"⚠️  [文本提取] openpyxl引擎失败，尝试其他引擎: {str(e)}")
            
            try:
                # 尝试xlrd引擎
                df = pd.read_excel(file_path, engine='xlrd', sheet_name=None)
                
                for sheet_name, sheet_data in df.items():
                    text_parts.append(f"=== 工作表: {sheet_name} ===")
                    
                    if not sheet_data.empty:
                        headers = [str(col) for col in sheet_data.columns]
                        text_parts.append("表头: " + " | ".join(headers))
                        
                        for _, row in sheet_data.iterrows():
                            row_text = " | ".join([str(val) if pd.notna(val) else "" for val in row])
                            text_parts.append(row_text)
                    
                    text_parts.append("")
            
            except Exception as e2:
                print(f"❌ [文本提取] 所有Excel引擎都失败: {str(e2)}")
                raise
        
        return "\n".join(text_parts)
    
    def _extract_from_word(self, file_path: Path) -> str:
        """从Word文档中提取文本（支持.doc和.docx格式）"""

        text_parts = []
        file_ext = file_path.suffix.lower()

        try:
            # 处理老式 .doc 格式
            if file_ext == '.doc':
                if not HAS_TEXTRACT:
                    raise ImportError(
                        "textract 未安装，无法读取老式Word(.doc)文档。\n"
                        "请运行: pip install textract\n"
                        "Windows还需安装antiword并加入PATH"
                    )

                print(f"📄 [文本提取] 检测到老式Word格式，使用textract提取")

                try:
                    text_bytes = textract.process(str(file_path.absolute()))
                    text = text_bytes.decode('utf-8', errors='ignore')
                except UnicodeDecodeError:
                    text = text_bytes.decode('gbk', errors='ignore')
                except Exception as e:
                    print(f"⚠️  [文本提取] textract提取失败: {str(e)}")
                    raise

                text = text.strip()
                return text

            # 处理新式 .docx 格式
            elif file_ext == '.docx':
                if Document is None:
                    raise ImportError("python-docx 未安装，无法读取Word文档")

                doc = Document(file_path)

                # 提取段落文本
                for para in doc.paragraphs:
                    if para.text.strip():
                        text_parts.append(para.text.strip())

                # 提取表格文本
                for table in doc.tables:
                    text_parts.append("=== 表格 ===")
                    for row in table.rows:
                        row_text = " | ".join([cell.text.strip() for cell in row.cells])
                        text_parts.append(row_text)
                    text_parts.append("")

                return "\n".join(text_parts)

        except Exception as e:
            print(f"❌ [文本提取] Word文档读取失败: {str(e)}")
            raise

        return "\n".join(text_parts)
    
    def _extract_from_ppt(self, file_path: Path) -> str:
        """从PowerPoint演示文稿中提取文本"""
        
        if Presentation is None:
            raise ImportError("python-pptx 未安装，无法读取PowerPoint文档")
        
        text_parts = []
        
        try:
            prs = Presentation(file_path)
            
            for slide_idx, slide in enumerate(prs.slides):
                text_parts.append(f"=== 幻灯片 {slide_idx + 1} ===")
                
                # 提取文本框内容
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        text_parts.append(shape.text.strip())
                    
                    # 提取表格内容
                    if shape.has_table:
                        for row in shape.table.rows:
                            row_text = " | ".join([cell.text.strip() for cell in row.cells])
                            text_parts.append(row_text)
                
                text_parts.append("")
        
        except Exception as e:
            print(f"❌ [文本提取] PowerPoint文档读取失败: {str(e)}")
            raise
        
        return "\n".join(text_parts)
    
    def clean_text(self, text: str) -> str:
        """清理提取的文本"""
        
        # 移除多余的空行
        text = re.sub(r'\n\s*\n\s*\n', '\n\n', text)
        
        # 移除行首尾空格
        lines = [line.strip() for line in text.split('\n')]
        text = '\n'.join(lines)
        
        return text


def main():
    """主函数"""
    if len(sys.argv) < 3:
        print("用法: python extract_air_compressor_data_llm.py <输入文件路径> <输出文本路径> [原始文件名]")
        sys.exit(1)
    
    input_file_path = sys.argv[1]
    output_text_path = sys.argv[2]
    original_filename = sys.argv[3] if len(sys.argv) > 3 else Path(input_file_path).name
    
    print(f"🚀 [文本提取] 开始处理文件")
    print(f"📁 [文本提取] 输入文件: {input_file_path}")
    print(f"📁 [文本提取] 输出路径: {output_text_path}")
    print(f"📁 [文本提取] 原始文件名: {original_filename}")
    
    try:
        # 创建文本提取器
        extractor = TextExtractor()
        
        # 提取文本
        extracted_text = extractor.extract_from_file(input_file_path)
        
        # 清理文本
        cleaned_text = extractor.clean_text(extracted_text)
        
        # 准备输出数据
        output_data = {
            "original_filename": original_filename,
            "file_type": extractor.file_type,
            "text_content": cleaned_text,
            "text_length": len(cleaned_text),
            "extraction_status": "success"
        }
        
        # 保存文本内容
        with open(output_text_path, 'w', encoding='utf-8') as f:
            json.dump(output_data, f, ensure_ascii=False, indent=2)
        
        print(f"✅ [文本提取] 处理完成")
        print(f"📄 结果已保存到: {output_text_path}")
        print(f"📊 提取文本长度: {len(cleaned_text)} 字符")
        
    except Exception as e:
        print(f"❌ [文本提取] 处理失败: {str(e)}")
        
        # 保存错误结果
        error_result = {
            "original_filename": original_filename,
            "file_type": Path(input_file_path).suffix.lower(),
            "text_content": "",
            "text_length": 0,
            "extraction_status": "failed",
            "error": str(e)
        }
        
        with open(output_text_path, 'w', encoding='utf-8') as f:
            json.dump(error_result, f, ensure_ascii=False, indent=2)
        
        sys.exit(1)


if __name__ == "__main__":
    main()