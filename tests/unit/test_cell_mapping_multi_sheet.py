"""测试 CellMapping 多sheet支持"""
import pytest
from config.cell_mapping import CellMapping


def test_get_sheet_name_upper():
    """测试获取上限sheet名称"""
    mapper = CellMapping()
    sheet_name = mapper.get_sheet_name("upper")
    assert sheet_name == "【1】按冷站电量数据-上限 "


def test_get_sheet_name_lower():
    """测试获取下限sheet名称"""
    mapper = CellMapping()
    sheet_name = mapper.get_sheet_name("lower")
    assert sheet_name == "【2】按冷站电量数据-下限"


def test_get_sheet_name_invalid():
    """测试无效sheet类型抛出异常"""
    mapper = CellMapping()
    with pytest.raises(ValueError, match="无效的sheet类型"):
        mapper.get_sheet_name("invalid")


def test_default_sheet_is_lower():
    """测试默认sheet是下限（向后兼容）"""
    mapper = CellMapping()
    assert mapper.SHEET_NAME == "【2】按冷站电量数据-下限"


def test_get_cell_works_for_both_sheets():
    """测试单元格映射对两个sheet都有效"""
    mapper = CellMapping()

    # 上限和下限使用相同的行列映射
    cell = mapper.get_cell("monthly_data", 0, "total_energy")
    assert cell == "D27"  # 1月总能耗位置


def test_sheet_names_exist_in_template():
    """集成测试：验证sheet名称在实际Excel模板中存在"""
    import openpyxl
    from pathlib import Path

    # 使用标准版模板（两个模板的sheet结构相同）
    template_path = Path("00_sample/模板测算表/节能评估表-按冷站电量数据-标准版.xlsx")

    if not template_path.exists():
        pytest.skip(f"模板文件不存在: {template_path}")

    wb = openpyxl.load_workbook(template_path)

    mapper = CellMapping()
    upper_sheet = mapper.get_sheet_name("upper")
    lower_sheet = mapper.get_sheet_name("lower")

    # 验证sheet名称在模板中存在
    assert upper_sheet in wb.sheetnames, (
        f"上限sheet不存在。期望: {repr(upper_sheet)}, "
        f"实际sheets: {wb.sheetnames}"
    )
    assert lower_sheet in wb.sheetnames, (
        f"下限sheet不存在。期望: {repr(lower_sheet)}, "
        f"实际sheets: {wb.sheetnames}"
    )

    wb.close()
