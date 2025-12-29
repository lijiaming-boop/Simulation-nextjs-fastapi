"""测试CellMapping配置与实际Excel模板的sheet名称匹配"""
import pytest
import openpyxl
from pathlib import Path
from config.cell_mapping import CellMapping


class TestSheetNameMatching:
    """验证配置的sheet名称与实际Excel模板匹配"""

    @pytest.fixture
    def excel_template(self):
        """获取Excel模板路径"""
        template_path = Path("00_sample/模板测算表/节能评估表-按冷站电量数据-标准版.xlsx")
        assert template_path.exists(), f"模板文件不存在: {template_path}"
        return template_path

    @pytest.fixture
    def mapper(self):
        """获取CellMapping实例"""
        return CellMapping()

    def test_upper_sheet_name_exists_in_template(self, excel_template, mapper):
        """测试上限sheet名称在Excel模板中存在"""
        # Given: 打开Excel模板
        wb = openpyxl.load_workbook(excel_template, read_only=True)
        actual_sheet_names = wb.sheetnames
        wb.close()

        # When: 获取配置的上限sheet名称
        configured_upper_name = mapper.get_sheet_name("upper")

        # Then: 配置的名称应该在实际sheet名称列表中
        assert configured_upper_name in actual_sheet_names, (
            f"配置的上限sheet名称 '{configured_upper_name}' 不在Excel模板中。\n"
            f"实际sheet名称: {actual_sheet_names}\n"
            f"提示: 检查是否有多余的空格"
        )

    def test_lower_sheet_name_exists_in_template(self, excel_template, mapper):
        """测试下限sheet名称在Excel模板中存在"""
        # Given: 打开Excel模板
        wb = openpyxl.load_workbook(excel_template, read_only=True)
        actual_sheet_names = wb.sheetnames
        wb.close()

        # When: 获取配置的下限sheet名称
        configured_lower_name = mapper.get_sheet_name("lower")

        # Then: 配置的名称应该在实际sheet名称列表中
        assert configured_lower_name in actual_sheet_names, (
            f"配置的下限sheet名称 '{configured_lower_name}' 不在Excel模板中。\n"
            f"实际sheet名称: {actual_sheet_names}"
        )

    def test_can_access_upper_sheet(self, excel_template, mapper):
        """测试可以使用配置的名称访问上限sheet"""
        # Given: 打开Excel模板
        wb = openpyxl.load_workbook(excel_template)
        configured_upper_name = mapper.get_sheet_name("upper")

        # When: 尝试访问上限sheet
        # Then: 应该不抛出KeyError
        try:
            sheet = wb[configured_upper_name]
            assert sheet is not None
        except KeyError as e:
            pytest.fail(
                f"无法访问上限sheet '{configured_upper_name}': {e}\n"
                f"可用的sheet名称: {wb.sheetnames}"
            )
        finally:
            wb.close()

    def test_can_access_lower_sheet(self, excel_template, mapper):
        """测试可以使用配置的名称访问下限sheet"""
        # Given: 打开Excel模板
        wb = openpyxl.load_workbook(excel_template)
        configured_lower_name = mapper.get_sheet_name("lower")

        # When: 尝试访问下限sheet
        # Then: 应该不抛出KeyError
        try:
            sheet = wb[configured_lower_name]
            assert sheet is not None
        finally:
            wb.close()
