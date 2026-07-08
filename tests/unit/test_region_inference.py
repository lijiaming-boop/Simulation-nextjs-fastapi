#!/usr/bin/env python3
"""测试地区推断和站房名称提取功能"""

import pytest
import pandas as pd
from pathlib import Path
import sys

# 添加项目根目录到 Python 路径
project_root = Path(__file__).parent.parent.parent
sys.path.insert(0, str(project_root / '.claude/skills/extract-and-check-data/scripts'))

from extract_and_check_data import extract_station_name, infer_region_from_text, extract_region_from_filename


class TestExtractStationName:
    """测试站房名称提取"""

    def test_extract_station_name_from_df(self):
        """测试从 DataFrame 提取站房名称"""
        # 模拟冷站 sheet 的前几行
        data = {
            0: [None, None, None],
            1: [None, '站房名称', None],
            2: [None, None, None],
            3: [None, 'F栋制冷站', None]
        }
        df = pd.DataFrame(data)

        result = extract_station_name(df)
        assert result == 'F栋制冷站'

    def test_extract_station_name_not_found(self):
        """测试站房名称字段不存在"""
        df = pd.DataFrame({
            0: [None, None],
            1: ['其他', '内容']
        })

        result = extract_station_name(df)
        assert result is None

    def test_extract_station_name_empty_value(self):
        """测试站房名称值为空"""
        data = {
            0: [None, None],
            1: [None, '站房名称'],
            2: [None, None],
            3: [None, None]
        }
        df = pd.DataFrame(data)

        result = extract_station_name(df)
        assert result is None


class TestInferRegionFromText:
    """测试从文本推断地区"""

    @pytest.mark.parametrize("text,expected_region", [
        # 站房名称示例
        ('北京总部冷站', '华北'),
        ('上海分公司机房', '华东'),
        ('深圳工业园区站房', '华南'),
        ('杭州研发中心空调机房', '华东'),
        ('武汉生产基地冷站', '华中'),
        ('成都办公楼制冷站', '西南'),
        ('西安厂区空调站', '西北'),
        ('沈阳冷站', '东北'),
        ('广州天河冷站', '华南'),
        ('石家庄仁和药业', '华北'),
        ('惠州光弘', '华南'),
        # 东南亚地区示例（返回具体城市名）
        ('吉隆坡艾杰旭汽车玻璃中央空调节能调研表', '吉隆坡'),
        ('曼谷某公司冷站', '曼谷'),
        ('胡志明市工业园制冷站', '胡志明市'),
        ('马尼拉中央空调机房', '马尼拉'),
        # 文件名示例
        ('【齐全】15.2%惠州光弘-中央空调节能调研表.xlsx', '华南'),
        ('【缺冬季】10.4%石家庄仁和药业中央空调节能调研表.xlsx', '华北'),
        ('【仅有夏季】6.5%和光同程-机房智能化调研表2023年5月版.xlsx', None),  # 无地区关键词
        ('上海和元-机房智能化调研表模板.xlsx', '华东'),
        ('北京某公司调研表.xlsx', '华北'),
        ('吉隆坡艾杰旭汽车玻璃中央空调节能调研表（首页含问题指引）-工业板块0131.xlsx', '吉隆坡'),
    ])
    def test_infer_region_with_city_keywords(self, text, expected_region):
        """测试包含城市关键词的文本"""
        result = infer_region_from_text(text)
        assert result == expected_region

    @pytest.mark.parametrize("text", [
        'F栋制冷站',
        '1号冷站',
        '中央空调机房',
        '某公司调研表.xlsx',
        None,
        '',
    ])
    def test_infer_region_no_keywords(self, text):
        """测试不包含地区关键词的文本"""
        result = infer_region_from_text(text)
        assert result is None


class TestExtractRegionFromFilename:
    """测试从文件名提取地区"""

    @pytest.mark.parametrize("filename,expected_region", [
        ('【齐全】15.2%惠州光弘-中央空调节能调研表.xlsx', '华南'),
        ('【缺冬季】10.4%石家庄仁和药业中央空调节能调研表.xlsx', '华北'),
        ('【仅有夏季】6.5%和光同程-机房智能化调研表2023年5月版.xlsx', None),
        ('【缺夏季】12.1%上海和元-机房智能化调研表模板（深度调研）.xlsx', '华东'),
        ('/path/to/北京某公司调研表.xlsx', '华北'),
        ('/path/to/深圳工业园区/调研表.xlsx', '华南'),
        ('test.xlsx', None),
    ])
    def test_extract_region_from_filename(self, filename, expected_region):
        """测试从文件名提取地区"""
        result = extract_region_from_filename(filename)
        assert result == expected_region


class TestExtractDataWithPriorityLogic:
    """测试完整的地区提取流程（三级优先级）"""

    def test_priority_1_station_name(self, tmp_path):
        """优先级1: 从站房名称提取地区"""
        import openpyxl

        wb = openpyxl.Workbook()

        # 创建"项目总体情况"sheet（不包含地区）
        ws_overview = wb.active
        ws_overview.title = '项目总体情况'
        ws_overview['A1'] = '项目名称'
        ws_overview['B1'] = '某公司冷站项目'

        # 创建"冷站"sheet（站房名称包含地区）
        ws_station = wb.create_sheet('冷站')
        ws_station['B2'] = '站房名称'
        ws_station['D2'] = '深圳工业园区F栋制冷站'

        # 添加设备运行规律标题行和月份行（模拟真实调研表结构）
        ws_station['C30'] = '设备运行规律'
        ws_station['C35'] = '月份'  # 月份标识
        months = ['1月', '2月', '3月', '4月', '5月', '6月', '7月', '8月', '9月', '10月', '11月', '12月']
        for i, month in enumerate(months):
            col_letter = chr(ord('D') + i)  # D, E, F, ..., O
            ws_station[f'{col_letter}35'] = month

        ws_station['C38'] = '冷机供水温度'
        ws_station['C39'] = '冷却泵运行频率'
        ws_station['C40'] = '冷却侧供回水温差'
        ws_station['C41'] = '冷冻泵运行频率'
        ws_station['C42'] = '冷冻侧供回水温差'
        ws_station['C44'] = '冷却塔出水温度'

        # 用电情况统计
        ws_station['C46'] = '用电情况统计'
        ws_station['C48'] = '月份'
        for i, month in enumerate(months):
            col_letter = chr(ord('D') + i)
            ws_station[f'{col_letter}48'] = month

        ws_station['C49'] = '整个冷站'  # 能耗标识
        # 只为前3个月添加数据，其他月份留空
        ws_station['D49'] = 10.5
        ws_station['E49'] = 12.3
        ws_station['F49'] = 15.8
        ws_station['C50'] = '万kWh'

        # 保存测试文件（文件名也不包含地区）
        test_file = tmp_path / 'test.xlsx'
        wb.save(test_file)

        # 执行提取
        from extract_and_check_data import extract_data_from_excel
        result = extract_data_from_excel(str(test_file))

        # 断言地区从站房名称提取成功
        assert result['project_region'] == '华南'

    def test_priority_2_filename(self, tmp_path):
        """优先级2: 站房名称失败，从文件名提取地区"""
        import openpyxl

        wb = openpyxl.Workbook()

        # 创建"项目总体情况"sheet（不包含地区）
        ws_overview = wb.active
        ws_overview.title = '项目总体情况'
        ws_overview['A1'] = '项目名称'
        ws_overview['B1'] = '某公司冷站项目'

        # 创建"冷站"sheet（站房名称不包含地区）
        ws_station = wb.create_sheet('冷站')
        ws_station['B2'] = '站房名称'
        ws_station['D2'] = 'F栋制冷站'

        # 添加设备运行规律标题行和月份行（模拟真实调研表结构）
        ws_station['C30'] = '设备运行规律'
        ws_station['C35'] = '月份'  # 月份标识
        months = ['1月', '2月', '3月', '4月', '5月', '6月', '7月', '8月', '9月', '10月', '11月', '12月']
        for i, month in enumerate(months):
            col_letter = chr(ord('D') + i)  # D, E, F, ..., O
            ws_station[f'{col_letter}35'] = month

        ws_station['C38'] = '冷机供水温度'
        ws_station['C39'] = '冷却泵运行频率'
        ws_station['C40'] = '冷却侧供回水温差'
        ws_station['C41'] = '冷冻泵运行频率'
        ws_station['C42'] = '冷冻侧供回水温差'
        ws_station['C44'] = '冷却塔出水温度'

        # 用电情况统计
        ws_station['C46'] = '用电情况统计'
        ws_station['C48'] = '月份'
        for i, month in enumerate(months):
            col_letter = chr(ord('D') + i)
            ws_station[f'{col_letter}48'] = month

        ws_station['C49'] = '整个冷站'  # 能耗标识
        # 只为前3个月添加数据，其他月份留空
        ws_station['D49'] = 10.5
        ws_station['E49'] = 12.3
        ws_station['F49'] = 15.8
        ws_station['C50'] = '万kWh'

        # 保存测试文件（文件名包含地区）
        test_file = tmp_path / '北京某公司调研表.xlsx'
        wb.save(test_file)

        # 执行提取
        from extract_and_check_data import extract_data_from_excel
        result = extract_data_from_excel(str(test_file))

        # 断言地区从文件名提取成功
        assert result['project_region'] == '华北'

    def test_priority_3_overview_sheet(self, tmp_path):
        """优先级3: 站房名称和文件名都失败，从项目总体情况提取"""
        import openpyxl

        wb = openpyxl.Workbook()

        # 创建"项目总体情况"sheet（包含地区）
        ws_overview = wb.active
        ws_overview.title = '项目总体情况'
        ws_overview['A1'] = '项目名称'
        ws_overview['B1'] = '武汉市某公司冷站项目'

        # 创建"冷站"sheet（站房名称不包含地区）
        ws_station = wb.create_sheet('冷站')
        ws_station['B2'] = '站房名称'
        ws_station['D2'] = 'F栋制冷站'

        # 添加设备运行规律标题行和月份行（模拟真实调研表结构）
        ws_station['C30'] = '设备运行规律'
        ws_station['C35'] = '月份'  # 月份标识
        months = ['1月', '2月', '3月', '4月', '5月', '6月', '7月', '8月', '9月', '10月', '11月', '12月']
        for i, month in enumerate(months):
            col_letter = chr(ord('D') + i)  # D, E, F, ..., O
            ws_station[f'{col_letter}35'] = month

        ws_station['C38'] = '冷机供水温度'
        ws_station['C39'] = '冷却泵运行频率'
        ws_station['C40'] = '冷却侧供回水温差'
        ws_station['C41'] = '冷冻泵运行频率'
        ws_station['C42'] = '冷冻侧供回水温差'
        ws_station['C44'] = '冷却塔出水温度'

        # 用电情况统计
        ws_station['C46'] = '用电情况统计'
        ws_station['C48'] = '月份'
        for i, month in enumerate(months):
            col_letter = chr(ord('D') + i)
            ws_station[f'{col_letter}48'] = month

        ws_station['C49'] = '整个冷站'  # 能耗标识
        # 只为前3个月添加数据，其他月份留空
        ws_station['D49'] = 10.5
        ws_station['E49'] = 12.3
        ws_station['F49'] = 15.8
        ws_station['C50'] = '万kWh'

        # 保存测试文件（文件名不包含地区）
        test_file = tmp_path / 'test.xlsx'
        wb.save(test_file)

        # 执行提取
        from extract_and_check_data import extract_data_from_excel
        result = extract_data_from_excel(str(test_file))

        # 断言地区从项目总体情况提取成功
        assert result['project_region'] == '华中'