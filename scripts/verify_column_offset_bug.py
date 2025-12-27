#!/usr/bin/env python3
"""验证月份列错位问题"""

import openpyxl
import pandas as pd

def verify_column_offset():
    """验证调研表中月份列的错位问题"""
    file_path = '00_sample/示例调研表/【齐全】15.2%惠州光弘-中央空调节能调研表.xlsx'

    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb['冷站']

    print("=== 设备运行规律（第32行 - 月份标题）===")
    print("月份列起始位置：")
    for col_idx in range(1, 15):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        value = ws[f'{col_letter}32'].value
        if value and '月' in str(value):
            print(f"  {col_letter}32 (索引{col_idx-1}): {value}")

    print("\n第34行（冷机供水温度）数据列：")
    for col_idx in range(3, 15):  # C列开始
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        value = ws[f'{col_letter}34'].value
        print(f"  {col_letter}34: {value}")

    print("\n=== 用电情况统计（第42行 - 月份标题）===")
    print("月份列起始位置：")
    for col_idx in range(1, 15):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        value = ws[f'{col_letter}42'].value
        if value and '月' in str(value):
            print(f"  {col_letter}42 (索引{col_idx-1}): {value}")

    print("\n第44行（冷站用电量）数据列：")
    for col_idx in range(2, 15):  # B列开始
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        value = ws[f'{col_letter}44'].value
        print(f"  {col_letter}44: {value}")

    print("\n=== 问题总结 ===")
    print("❌ 设备运行规律：1月数据在 D 列（索引3）")
    print("❌ 用电情况统计：1月数据在 C 列（索引2）")
    print("❌ 当前代码使用 month_col_start=3（D列）会导致用电数据错位：")
    print("   - 漏读1月数据（C列）")
    print("   - 2月数据（D列）被错当成1月")
    print("   - 后续月份全部向左错位一个月")

if __name__ == '__main__':
    verify_column_offset()
