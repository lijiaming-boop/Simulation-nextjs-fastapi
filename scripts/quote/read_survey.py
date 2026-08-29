#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
读取标准能管项目调研表，输出结构化 JSON。
按照 SKILL.md 中的单元格映射规则提取数据。
"""

import argparse
import json
import os
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)

# ========== 默认配置（基于脚本位置，避免绝对路径） ==========
PROJECT_ROOT = Path(__file__).parent.parent.parent
DEFAULT_SURVEY_PATH = str(PROJECT_ROOT / "00_sample" / "报价模板" / "标准能管项目调研表_V1.4_已填写.xlsx")
DEFAULT_OUTPUT_PATH = str(PROJECT_ROOT / "outputs" / "quote" / "test_ai_project.json")

def get_cell_value(ws, row, col_letter):
    """安全读取单元格值，返回 None 如果单元格为空"""
    cell = ws[f"{col_letter}{row}"]
    val = cell.value
    if val is None:
        return None
    if isinstance(val, str):
        val = val.strip()
        if val == "" or val == "-" or val == "/":
            return None
    return val

def get_cell_int(ws, row, col_letter):
    """读取整数值，非数字返回 0"""
    val = get_cell_value(ws, row, col_letter)
    if val is None:
        return 0
    try:
        return int(val)
    except (ValueError, TypeError):
        return 0

def get_cell_bool(ws, row, col_letter):
    """读取布尔值：'选用'/'是' → True"""
    val = get_cell_value(ws, row, col_letter)
    if val is None:
        return False
    s = str(val).strip()
    return s in ("选用", "是", "TRUE", "True", "true", "yes", "Y", "1")

def get_cell_float(ws, row, col_letter):
    """读取浮点数值"""
    val = get_cell_value(ws, row, col_letter)
    if val is None:
        return 0.0
    try:
        return float(val)
    except (ValueError, TypeError):
        return 0.0

def map_deployment_mode(raw):
    """映射部署方式"""
    if raw is None:
        return "cloud"
    s = str(raw).strip()
    mapping = {
        "云端部署": "cloud",
        "私有部署（我方服务器/工控机）": "private_our",
        "私有部署（客户服务器/工控机）": "private_client",
    }
    # 模糊匹配
    if "云端" in s:
        return "cloud"
    if "我方" in s:
        return "private_our"
    if "客户" in s:
        return "private_client"
    return s  # fallback

def main():
    parser = argparse.ArgumentParser(description="读取标准能管项目调研表，输出结构化 JSON")
    parser.add_argument("--survey", "-s", default=DEFAULT_SURVEY_PATH, help="输入调研表 Excel 路径")
    parser.add_argument("--output", "-o", default=DEFAULT_OUTPUT_PATH, help="输出 JSON 路径")
    args = parser.parse_args()

    survey_path = args.survey
    output_path = args.output

    print(f"📂 读取调研表: {survey_path}")
    wb = openpyxl.load_workbook(survey_path, data_only=True)
    ws = wb.active
    print(f"   Sheet: {ws.title}")
    print(f"   最大行: {ws.max_row}, 最大列: {ws.max_column}")

    # ========== 第一步：打印所有非空单元格 ==========
    print("\n" + "="*70)
    print("📋 所有非空单元格:")
    print("="*70)
    non_empty_count = 0
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            if cell.value is not None:
                val = str(cell.value).strip()
                if val:
                    print(f"   {cell.coordinate:>6} = {repr(cell.value)}")
                    non_empty_count += 1
    print(f"\n   共 {non_empty_count} 个非空单元格")

    # ========== 第二步：按映射规则提取 ==========
    print("\n" + "="*70)
    print("📊 按 SKILL.md 映射规则提取数据")
    print("="*70)

    # --- 板块一：项目基础信息 ---
    client_name = get_cell_value(ws, 7, "C") or ""
    project_name = get_cell_value(ws, 8, "C") or ""
    project_location = get_cell_value(ws, 9, "C") or ""
    deployment_mode_raw = get_cell_value(ws, 12, "D")
    deployment_mode = map_deployment_mode(deployment_mode_raw)
    energy_mgmt_method = get_cell_value(ws, 13, "D") or ""

    print(f"\n🏢 板块一：项目基础信息")
    print(f"   client_name: {client_name}")
    print(f"   project_name: {project_name}")
    print(f"   project_location: {project_location}")
    print(f"   deployment_mode (raw): {deployment_mode_raw} → {deployment_mode}")
    print(f"   energy_mgmt_method (D13): {energy_mgmt_method}")

    # --- 板块二：既有表计 ---
    electric_accessible = get_cell_int(ws, 18, "C")
    electric_need_retrofit = get_cell_int(ws, 18, "D")
    water_accessible = get_cell_int(ws, 19, "C")
    water_need_retrofit = get_cell_int(ws, 19, "D")

    print(f"\n📟 板块二：既有表计")
    print(f"   电表 可接入={electric_accessible}, 需改造={electric_need_retrofit}")
    print(f"   水表 可接入={water_accessible}, 需改造={water_need_retrofit}")

    # --- 板块三：新增电表 (rows 29-34) ---
    electric_meters = []
    for row in range(29, 35):
        qty = get_cell_value(ws, row, "C")
        if qty is not None:
            try:
                qty_int = int(qty)
            except (ValueError, TypeError):
                qty_int = 0
            if qty_int > 0:
                loc = get_cell_value(ws, row, "B") or ""
                ct_current = get_cell_int(ws, row, "D")
                ct_type = get_cell_value(ws, row, "E") or ""
                site_cond = get_cell_value(ws, row, "F") or ""
                electric_meters.append({
                    "location": loc,
                    "quantity": qty_int,
                    "ct_current": ct_current,
                    "ct_type": ct_type,
                    "site_condition": site_cond
                })
    print(f"\n⚡ 板块三：新增电表 ({len(electric_meters)} 行)")
    for m in electric_meters:
        print(f"   {m}")

    # --- 板块四：新增水表 (rows 39-46) ---
    water_meters = []
    for row in range(39, 47):
        qty = get_cell_value(ws, row, "D")
        if qty is not None:
            try:
                qty_int = int(qty)
            except (ValueError, TypeError):
                qty_int = 0
            if qty_int > 0:
                loc = get_cell_value(ws, row, "B") or ""
                dn = get_cell_value(ws, row, "C") or ""
                need_install = get_cell_bool(ws, row, "E")
                site_cond = get_cell_value(ws, row, "F") or ""
                water_meters.append({
                    "location": loc,
                    "dn": dn,
                    "quantity": qty_int,
                    "need_install": need_install,
                    "site_condition": site_cond
                })
    print(f"\n💧 板块四：新增水表 ({len(water_meters)} 行)")
    for m in water_meters:
        print(f"   {m}")

    # --- 板块五：新增流量计/冷热量表 (rows 49-56) ---
    flow_meters = []
    for row in range(49, 57):
        qty = get_cell_value(ws, row, "D")
        if qty is not None:
            try:
                qty_int = int(qty)
            except (ValueError, TypeError):
                qty_int = 0
            if qty_int > 0:
                medium = get_cell_value(ws, row, "B") or ""
                dn = get_cell_value(ws, row, "C") or ""
                install_type = get_cell_value(ws, row, "E") or ""
                site_cond = get_cell_value(ws, row, "F") or ""
                flow_meters.append({
                    "medium": medium,
                    "dn": dn,
                    "quantity": qty_int,
                    "install_type": install_type,
                    "site_condition": site_cond
                })
    print(f"\n🌡️ 板块五：新增流量计/冷热量表 ({len(flow_meters)} 行)")
    for m in flow_meters:
        print(f"   {m}")

    # --- 板块六：第三方系统 (rows 59-64) ---
    third_party = []
    for row in range(59, 65):
        name = get_cell_value(ws, row, "B")
        if name is not None:
            sys_type = get_cell_value(ws, row, "C") or ""
            direction = get_cell_value(ws, row, "D") or ""
            desc = get_cell_value(ws, row, "E") or ""
            third_party.append({
                "name": name,
                "type": sys_type,
                "direction": direction,
                "description": desc
            })
    print(f"\n🔗 板块六：第三方系统 ({len(third_party)} 行)")
    for m in third_party:
        print(f"   {m}")

    # --- 板块七：可选软件 (rows 69-75) ---
    selected_modules = []
    for r in range(69, 76):
        opt_name = get_cell_value(ws, r, "B")
        if not opt_name:
            continue
        opt_name = str(opt_name).strip()
        c_val = get_cell_value(ws, r, "C")
        c_text = str(c_val).strip() if c_val else ""
        scope = str(get_cell_value(ws, r, "D") or "").strip()
        qty_name = str(get_cell_value(ws, r, "E") or "").strip()
        cust_req = str(get_cell_value(ws, r, "F") or "").strip()
        explicitly_selected = "选用" in c_text and "不选" not in c_text
        has_content = any(
            v and v not in ("-", "无", "0")
            for v in (scope, qty_name, cust_req)
        )
        if not explicitly_selected and not has_content:
            continue
        selected_modules.append({
            "name": opt_name,
            "scope": scope,
            "quantity_name": qty_name,
            "customer_request": cust_req,
        })

    print(f"\n📱 板块七：可选软件 ({len(selected_modules)} 项选用)")
    for m in selected_modules:
        print(f"   {m['name']} | 范围:{m['scope']} | 数量:{m['quantity_name']}")

    # --- 板块七-附加：大屏/组态/报表 (rows 79-84) ---
    extra_dashboards = get_cell_int(ws, 80, "C")
    extra_topology = get_cell_int(ws, 81, "C")
    extra_reports = get_cell_int(ws, 82, "C")

    print(f"\n🖥️ 板块七-附加：额外大屏/组态/报表")
    print(f"   额外大屏: {extra_dashboards}")
    print(f"   额外组态: {extra_topology}")
    print(f"   额外报表: {extra_reports}")

    # --- 板块八：项目需求补充 (rows 89-91) ---
    current_status = get_cell_value(ws, 89, "C") or ""
    problems = get_cell_value(ws, 90, "C") or ""
    goals = get_cell_value(ws, 91, "C") or ""

    print(f"\n📝 板块八：项目需求补充")
    print(f"   current_status: {current_status}")
    print(f"   problems: {problems}")
    print(f"   goals: {goals}")

    # ========== 第三步：构建 JSON ==========
    project = {
        "client_name": client_name,
        "project_name": project_name,
        "project_location": project_location,
        "deployment_mode": deployment_mode,
        "energy_mgmt_method": energy_mgmt_method,
        "existing_meters": {
            "electric_accessible": electric_accessible,
            "electric_need_retrofit": electric_need_retrofit,
            "water_accessible": water_accessible,
            "water_need_retrofit": water_need_retrofit
        },
        "new_meters": {
            "electric": electric_meters,
            "water": water_meters,
            "flow_meters": flow_meters
        },
        "third_party_systems": third_party,
        "optional_modules": {
            "selected_modules": selected_modules,
            "extra_dashboards": extra_dashboards,
            "extra_topology": extra_topology,
            "extra_reports": extra_reports
        },
        "project_notes": {
            "current_status": current_status,
            "problems": problems,
            "goals": goals
        }
    }

    # ========== 第四步：写入 JSON ==========
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(project, f, ensure_ascii=False, indent=2)

    print(f"\n✅ JSON 已写入: {output_path}")

    # ========== 第五步：摘要 ==========
    print("\n" + "="*70)
    print("📊 提取摘要")
    print("="*70)
    print(f"   客户名称: {client_name}")
    print(f"   项目名称: {project_name}")
    print(f"   项目地点: {project_location}")
    print(f"   部署方式: {deployment_mode}")
    print(f"   能源管理方式: {energy_mgmt_method}")
    print(f"   既有电表: 可接入 {electric_accessible} / 需改造 {electric_need_retrofit}")
    print(f"   既有水表: 可接入 {water_accessible} / 需改造 {water_need_retrofit}")
    print(f"   新增电表: {len(electric_meters)} 项, 合计 {sum(m['quantity'] for m in electric_meters)} 台")
    print(f"   新增水表: {len(water_meters)} 项, 合计 {sum(m['quantity'] for m in water_meters)} 台")
    print(f"   新增流量计/冷热量表: {len(flow_meters)} 项, 合计 {sum(m['quantity'] for m in flow_meters)} 台")
    print(f"   第三方系统: {len(third_party)} 个")
    print(f"   可选模块: {len(selected_modules)} 项选用")
    for m in selected_modules:
        print(f"     - {m['name']}")
    print(f"   额外: 大屏={extra_dashboards}, 组态={extra_topology}, 报表={extra_reports}")
    print(f"\n🎉 完成!")

if __name__ == "__main__":
    main()
