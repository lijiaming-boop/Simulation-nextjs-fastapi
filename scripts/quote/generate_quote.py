"""
标准能管报价表生成器
从 internal_bottom_sheet.json 生成标准价版和渠道价版报价 xlsx
格式参考：报价.xlsx（渠道价）、2025年能管大公辅软件报价-202509.xlsx（标准价）
"""
import json
import sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill


THIN_BORDER = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

TITLE_FONT = Font(name='宋体', size=16, bold=True)
HEADER_FONT = Font(name='宋体', size=12, bold=True)
HEADER_FONT_SM = Font(name='宋体', size=10, bold=True)
NORMAL_FONT = Font(name='宋体', size=10)
SECTION_FONT = Font(name='宋体', size=10, bold=True)
TOTAL_FONT = Font(name='宋体', size=10, bold=True)
GRAND_TOTAL_FONT = Font(name='宋体', size=12, bold=True)
NOTE_FONT = Font(name='宋体', size=10)

HEADER_FILL = PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid')
SECTION_FILL = PatternFill(start_color='E8E6E7', end_color='E8E6E7', fill_type='solid')
SUBTOTAL_FILL = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
TOTAL_FILL = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')

CENTER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT_ALIGN = Alignment(horizontal='left', vertical='center', wrap_text=True)
RIGHT_ALIGN = Alignment(horizontal='right', vertical='center')
NOTE_ALIGN = Alignment(horizontal='left', vertical='top', wrap_text=True)

PRICE_FMT = '0.00_ '

COL_WIDTHS = {
    'A': 6.17, 'B': 27.22, 'C': 7.34, 'D': 7.50,
    'E': 82.86, 'F': 10.0, 'G': 22.0, 'H': 14.83,
    'I': 16.0, 'J': 21.11,
}

ZONES = [
    ('一', '平台软件', 'software_items'),
    ('二', '硬件设备及仪表', 'hardware_items'),
    ('三', '施工及调试', 'construction_items'),
]


def write_quote_xlsx(data: dict, price_mode: str, output_path: str):
    price_key = 'standard_price' if price_mode == 'standard' else 'channel_price'
    wb = Workbook()
    ws = wb.active
    ws.title = '报价模板'

    for col_letter, width in COL_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    row = 1
    sequence = [0]

    def next_seq():
        sequence[0] += 1
        return sequence[0]

    def apply_borders(r, col_start=1, col_end=10):
        for c in range(col_start, col_end + 1):
            ws.cell(row=r, column=c).border = THIN_BORDER

    def write_data_row(r, values, font=NORMAL_FONT, fill=None):
        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=r, column=col_idx, value=val)
            cell.border = THIN_BORDER
            cell.font = font
            if fill:
                cell.fill = fill
            if col_idx in (1, 3, 4):
                cell.alignment = CENTER_ALIGN
            elif col_idx in (8, 9):
                cell.alignment = RIGHT_ALIGN
            elif col_idx == 5:
                cell.alignment = LEFT_ALIGN
            else:
                cell.alignment = CENTER_ALIGN

    # ---- 标题 ----
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
    title_cell = ws.cell(row=row, column=1, value='云智控产品报价表')
    title_cell.font = TITLE_FONT
    title_cell.alignment = CENTER_ALIGN
    apply_borders(row)
    ws.row_dimensions[row].height = 38.85
    row += 1

    # ---- 表头 ----
    if price_mode == 'channel':
        price_header = '渠道单价（元）'
        total_header = '渠道总价（元）'
    else:
        price_header = '单价（元）'
        total_header = '总价（元）'
    headers = ['序号', '模块', '数量', '单位', '功能描述', '品牌', '型号',
               price_header, total_header, '备注']
    for col_idx, val in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col_idx, value=val)
        cell.font = HEADER_FONT_SM if col_idx == 5 else HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER
    ws.row_dimensions[row].height = 27.1
    row += 1

    # ---- 逐分区写入 ----
    subtotal_rows = []

    for zone_num, zone_name, item_key in ZONES:
        items = data.get(item_key, [])
        if not items:
            continue

        # 分区标题行：A=序号，B:J=名称（合并）
        ws.cell(row=row, column=1, value=zone_num).font = SECTION_FONT
        ws.cell(row=row, column=1).alignment = CENTER_ALIGN
        ws.cell(row=row, column=1).border = THIN_BORDER
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
        name_cell = ws.cell(row=row, column=2, value=zone_name)
        name_cell.font = SECTION_FONT
        name_cell.alignment = LEFT_ALIGN
        for c in range(1, 11):
            ws.cell(row=row, column=c).fill = SECTION_FILL
            ws.cell(row=row, column=c).border = THIN_BORDER
        ws.row_dimensions[row].height = 26.2
        row += 1

        # 数据行
        data_start_row = row
        for item in items:
            qty = item.get('quantity', 0)
            price = item.get(price_key, 0) or 0
            total = qty * price
            seq = next_seq()

            write_data_row(row, [
                seq,
                item.get('module', ''),
                qty,
                item.get('unit', ''),
                item.get('description', ''),
                item.get('brand', ''),
                item.get('model', ''),
                price if price else '',
                total if price else '',
                ''
            ])
            for c in (8, 9):
                if price:
                    ws.cell(row=row, column=c).number_format = PRICE_FMT
            row += 1
        data_end_row = row - 1

        # 小计行：B="小计"，C:H合并空，I=SUM公式
        ws.cell(row=row, column=1, value='').border = THIN_BORDER
        ws.cell(row=row, column=2, value='小计').font = TOTAL_FONT
        ws.cell(row=row, column=2).alignment = CENTER_ALIGN
        ws.cell(row=row, column=2).border = THIN_BORDER
        ws.cell(row=row, column=2).fill = SUBTOTAL_FILL
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=8)
        for c in range(3, 9):
            ws.cell(row=row, column=c).border = THIN_BORDER
            ws.cell(row=row, column=c).fill = SUBTOTAL_FILL
        sub_cell = ws.cell(row=row, column=9)
        sub_cell.border = THIN_BORDER
        sub_cell.fill = SUBTOTAL_FILL
        sub_cell.alignment = RIGHT_ALIGN
        sub_cell.font = TOTAL_FONT
        sub_cell.number_format = PRICE_FMT
        if data_end_row >= data_start_row:
            sub_cell.value = f'=SUM(I{data_start_row}:I{data_end_row})'
        subtotal_rows.append(row)
        row += 1
        row += 1  # 分区间距

    # ---- 税费分区 ----
    tax_rates = data.get('tax_breakdown', {}).get('rates', {})
    sw_hw_rate = tax_rates.get('software_hardware', 0.13)
    ct_rate = tax_rates.get('construction', 0.06)
    sw_hw_pct = int(sw_hw_rate * 100)
    ct_pct = int(ct_rate * 100)

    ws.cell(row=row, column=1, value='四').font = SECTION_FONT
    ws.cell(row=row, column=1).alignment = CENTER_ALIGN
    ws.cell(row=row, column=1).border = THIN_BORDER
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
    tax_title = f'税费（软硬件{sw_hw_pct}%，调试{ct_pct}%）'
    name_cell = ws.cell(row=row, column=2, value=tax_title)
    name_cell.font = SECTION_FONT
    name_cell.alignment = LEFT_ALIGN
    for c in range(1, 11):
        ws.cell(row=row, column=c).fill = SECTION_FILL
        ws.cell(row=row, column=c).border = THIN_BORDER
    ws.row_dimensions[row].height = 26.2
    row += 1

    # 增值税行
    tax_seq = next_seq()
    tax_desc = f'软硬件{sw_hw_pct}%，调试{ct_pct}%'
    write_data_row(row, [
        tax_seq, '增值税', 1, '项', tax_desc,
        '', '', '', '', ''
    ])

    # 税费公式：引用各分区小计行
    if len(subtotal_rows) >= 2:
        sw_sub = subtotal_rows[0]
        hw_sub = subtotal_rows[1] if len(subtotal_rows) > 1 else sw_sub
        ct_sub = subtotal_rows[2] if len(subtotal_rows) > 2 else None
        sw_hw_sub_total = f'I{sw_sub}+I{hw_sub}'
        if ct_sub:
            formula = f'={sw_hw_sub_total}*{sw_hw_rate}+I{ct_sub}*{ct_rate}'
        else:
            formula = f'={sw_hw_sub_total}*{sw_hw_rate}'
        ws.cell(row=row, column=9).value = formula
    else:
        sw_hw_total = sum(
            (i.get(price_key, 0) or 0) * i.get('quantity', 0)
            for i in data.get('software_items', []) + data.get('hardware_items', [])
        )
        const_total = sum(
            (i.get(price_key, 0) or 0) * i.get('quantity', 0)
            for i in data.get('construction_items', [])
        )
        ws.cell(row=row, column=9).value = round(
            sw_hw_total * sw_hw_rate + const_total * ct_rate, 2
        )
    ws.cell(row=row, column=9).number_format = PRICE_FMT
    ws.cell(row=row, column=9).alignment = RIGHT_ALIGN
    tax_row = row
    row += 1

    # 税费小计
    ws.cell(row=row, column=2, value='小计').font = TOTAL_FONT
    ws.cell(row=row, column=2).alignment = CENTER_ALIGN
    ws.cell(row=row, column=2).border = THIN_BORDER
    ws.cell(row=row, column=2).fill = SUBTOTAL_FILL
    ws.cell(row=row, column=9).value = f'=I{tax_row}'
    ws.cell(row=row, column=9).number_format = PRICE_FMT
    ws.cell(row=row, column=9).font = TOTAL_FONT
    ws.cell(row=row, column=9).fill = SUBTOTAL_FILL
    ws.cell(row=row, column=9).border = THIN_BORDER
    ws.cell(row=row, column=9).alignment = RIGHT_ALIGN
    subtotal_rows.append(row)
    row += 1

    # ---- 总计 ----
    ws.cell(row=row, column=1, value='').border = THIN_BORDER
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
    total_label = ws.cell(row=row, column=2, value='   总计')
    total_label.font = GRAND_TOTAL_FONT
    total_label.alignment = LEFT_ALIGN
    total_label.fill = TOTAL_FILL
    for c in range(2, 9):
        ws.cell(row=row, column=c).fill = TOTAL_FILL
        ws.cell(row=row, column=c).border = THIN_BORDER
    grand_cell = ws.cell(row=row, column=9)
    grand_cell.font = GRAND_TOTAL_FONT
    grand_cell.fill = TOTAL_FILL
    grand_cell.border = THIN_BORDER
    grand_cell.alignment = RIGHT_ALIGN
    grand_cell.number_format = PRICE_FMT
    sub_refs = '+'.join(f'I{r}' for r in subtotal_rows[:-1])
    grand_cell.value = f'={sub_refs}+I{tax_row}'
    row += 2

    # ---- 备注 ----
    note_start = row
    note_end = row + 3
    ws.merge_cells(start_row=note_start, start_column=1,
                   end_row=note_end, end_column=13)
    notes = [
        '备注：',
        '1、以上各模块价格均为买断模式，前三年不收取云资源及软件升级服务费用；第四年开始，每年按照云智控软件总价的10%的标准收取云资源及软件升级服务费；',
        f'2、云智控系统软件和硬件费含{sw_hw_pct}%的增值税，技术实施费、云资源及软件升级费含{ct_pct}%的增值税；',
        '3、质保政策：若为公有云部署，蘑菇物联品牌硬件3年质保，三方品硬件1年质保；若为私有云部署，所有硬件均1年质保。',
    ]
    note_cell = ws.cell(row=note_start, column=1, value='\n'.join(notes))
    note_cell.font = NOTE_FONT
    note_cell.alignment = NOTE_ALIGN
    ws.row_dimensions[note_start].height = 90

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    sw_hw_total = sum(
        (i.get(price_key, 0) or 0) * i.get('quantity', 0)
        for i in data.get('software_items', []) + data.get('hardware_items', [])
    )
    const_total = sum(
        (i.get(price_key, 0) or 0) * i.get('quantity', 0)
        for i in data.get('construction_items', [])
    )
    tax_amount = sw_hw_total * sw_hw_rate + const_total * ct_rate
    return sw_hw_total + const_total + tax_amount


def generate(bottom_sheet_path: str, output_dir: str):
    data = json.loads(Path(bottom_sheet_path).read_text(encoding='utf-8'))
    out_dir = Path(output_dir)

    std_total = write_quote_xlsx(data, 'standard', str(out_dir / '报价_标准价.xlsx'))
    ch_total = write_quote_xlsx(data, 'channel', str(out_dir / '报价_渠道价.xlsx'))

    return {
        'standard_grand_total': round(std_total, 2),
        'channel_grand_total': round(ch_total, 2),
    }


if __name__ == '__main__':
    bottom_path = sys.argv[1] if len(sys.argv) > 1 else 'outputs/quote/internal_bottom_sheet.json'
    out_dir = sys.argv[2] if len(sys.argv) > 2 else 'outputs/quote'

    result = generate(bottom_path, out_dir)
    print(f'[OK] 标准价报价: outputs/quote/报价_标准价.xlsx  总计: {result["standard_grand_total"]:,.2f}')
    print(f'[OK] 渠道价报价: outputs/quote/报价_渠道价.xlsx  总计: {result["channel_grand_total"]:,.2f}')
