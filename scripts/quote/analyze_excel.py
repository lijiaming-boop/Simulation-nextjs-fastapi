"""Analyze the pricing Excel structure, focusing on optional software modules."""
import openpyxl
import sys
import io

EXCEL_PATH = r"C:\Users\30937\Desktop\2025年云智控报价体系-全产品线11.xlsx"
OUTPUT_FILE = r"C:\Users\30937\Desktop\standard-energy-quote - 副本 (2) - 副本\scripts\excel_analysis_output.txt"

def safe_str(v):
    """Convert value to string, replacing non-breaking spaces."""
    if v is None:
        return ""
    return str(v).replace('\xa0', ' ')

def main():
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    out = io.StringIO()

    def p(s=""):
        out.write(s + "\n")

    p("=" * 80)
    p("SHEET NAMES:")
    p("=" * 80)
    for i, name in enumerate(wb.sheetnames):
        ws = wb[name]
        p(f"  [{i}] '{name}'  (rows={ws.max_row}, cols={ws.max_column})")

    # For each sheet, print header rows + first ~55 data rows
    for name in wb.sheetnames:
        ws = wb[name]
        p()
        p("=" * 80)
        p(f"SHEET: '{name}'  (rows={ws.max_row}, cols={ws.max_column})")
        p("=" * 80)

        max_show = min(ws.max_row, 55)
        for r in range(1, max_show + 1):
            row_vals = []
            for c in range(1, min(ws.max_column + 1, 20)):
                v = ws.cell(row=r, column=c).value
                if v is not None:
                    row_vals.append(f"C{c}={safe_str(v)}")
            if row_vals:
                p(f"  R{r:3d}: {' | '.join(row_vals)}")
            else:
                p(f"  R{r:3d}: (empty)")

        if ws.max_row > max_show:
            p(f"  ... ({ws.max_row - max_show} more rows omitted)")

    # ── FOCUS: Software sheets - show ALL rows for software sheets ──
    software_sheets = [n for n in wb.sheetnames if "软件" in n]
    if software_sheets:
        p()
        p("=" * 80)
        p("DETAILED: ALL SOFTWARE SHEET ROWS")
        p("=" * 80)
        for name in software_sheets:
            ws = wb[name]
            p(f"\n--- SHEET: '{name}' (ALL {ws.max_row} rows) ---")
            for r in range(1, ws.max_row + 1):
                row_vals = []
                for c in range(1, min(ws.max_column + 1, 20)):
                    v = ws.cell(row=r, column=c).value
                    if v is not None:
                        row_vals.append(f"C{c}={safe_str(v)}")
                if row_vals:
                    p(f"  R{r:3d}: {' | '.join(row_vals)}")

    # ── FOCUS: Identify 标配/选配 markers ──
    p()
    p("=" * 80)
    p("选配/标配 ANALYSIS: Rows containing '选配' or '标配'")
    p("=" * 80)
    for name in wb.sheetnames:
        ws = wb[name]
        found = []
        for r in range(1, ws.max_row + 1):
            for c in range(1, min(ws.max_column + 1, 20)):
                v = ws.cell(row=r, column=c).value
                if v and ("选配" in str(v) or "标配" in str(v)):
                    row_vals = []
                    for cc in range(1, min(ws.max_column + 1, 20)):
                        vv = ws.cell(row=r, column=cc).value
                        if vv is not None:
                            row_vals.append(f"C{cc}={safe_str(vv)}")
                    found.append((r, row_vals))
                    break
        if found:
            p(f"\n  SHEET '{name}': {len(found)} rows with 选配/标配")
            for r, vals in found:
                p(f"    R{r:3d}: {' | '.join(vals)}")

    wb.close()

    # Write output to file
    with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
        f.write(out.getvalue())
    print(f"Output written to: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
