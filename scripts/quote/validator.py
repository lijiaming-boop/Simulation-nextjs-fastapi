"""
标准能管输出校验器
检查报价 xlsx 和方案 docx 中的禁词、一致性等问题
"""
import json
import re
import sys
from pathlib import Path
from openpyxl import load_workbook
from docx import Document


# ===================== 禁词清单 =====================

FORBIDDEN_WORDS_QUOTE = [
    # 内部逻辑痕迹
    '点位取整', '点位向上', '点位公式', '云盒公式', '内部来源',
    '三方系统+1000', '三方系统+', '+1000点', '向上取整',
    '原始点位', '报价点位', 'CEILING',
    # 来源/映射
    '来源工作表', '报价体系内部', '内部映射', '内部底稿',
    'price_rows', 'source_materials',
    # 用户确认痕迹
    '用户确认', '待确认', '压测', '虚拟', '调研表',
    '内部风险', '脚本说明', 'AUTO', 'CONFIRM', 'INQUIRY',
    'NONSTANDARD', 'BLOCKED',
    # AI 相关
    'AI', '人工智能', '大模型', '机器学习', '深度学习',
    # 不应出现的价格信息
    '渠道底价', '标准底价', '联营价',
]

FORBIDDEN_WORDS_SOLUTION = [
    # AI / 高级分析
    'AI', '人工智能', '大模型', '机器学习', '深度学习',
    '能源AI', 'AI预测', 'AI自动分析', 'AI可视化',
    # 分期
    '一期', '二期', '一期功能', '二期功能', '二期实现',
    '二期可选', '第一年', '第二年', '分阶段实施',
    # 控制
    '自动启停', '自动优化', '闭环控制',
    # 承诺
    '固定节能率', '节能率', '诊断报告', '优化建议',
    '自动生成报告', '自动诊断',
    # 内部痕迹
    '压测', '虚拟项目', '调研表', '内部底稿',
    '默认模块', '报价映射', '点位取整',
    # 越界表达
    '能耗预测', '能源预测', '预测分析',  # 标准能管不做预测
]

# 方案中允许的例外词（在特定上下文中可以出现）
SOLUTION_ALLOWLIST = [
    # '节能' 本身不是问题，只有"固定节能率"才是
]

# 检查时忽略大小写
CASE_INSENSITIVE = [
    'AI', 'ai', 'Ai',
]


def extract_text_from_xlsx(filepath: str) -> str:
    """提取 xlsx 中所有文本"""
    wb = load_workbook(filepath, data_only=True)
    all_text = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell is not None:
                    all_text.append(str(cell))
    return ' '.join(all_text)


def extract_text_from_docx(filepath: str) -> str:
    """提取 docx 中所有文本"""
    doc = Document(filepath)
    all_text = []
    for para in doc.paragraphs:
        if para.text.strip():
            all_text.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    all_text.append(cell.text)
    return '\n'.join(all_text)


def extract_modules_from_xlsx(filepath: str) -> list:
    """从报价 xlsx 提取模块列表"""
    wb = load_workbook(filepath, data_only=True)
    modules = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=1, values_only=True):
            # 列B是模块名
            module = str(row[1]).strip() if len(row) > 1 and row[1] else ''
            qty = row[2] if len(row) > 2 else None
            if module and module not in ('模块', '小计', '总计', '') and not module.startswith('一') and not module.startswith('二') and not module.startswith('三') and not module.startswith('四'):
                try:
                    qty_val = float(qty) if qty else 0
                except (ValueError, TypeError):
                    qty_val = 0
                modules.append({'module': module, 'quantity': qty_val})
    return modules


def check_forbidden_words(text: str, forbidden_list: list, source_name: str) -> list:
    """检查禁词，返回命中的词列表"""
    hits = []
    text_lower = text.lower()
    for word in forbidden_list:
        if word.lower() in text_lower:
            # 找到上下文
            idx = text_lower.find(word.lower())
            start = max(0, idx - 30)
            end = min(len(text), idx + len(word) + 30)
            context = text[start:end].replace('\n', ' ')
            hits.append({
                'word': word,
                'context': f'...{context}...',
                'source': source_name,
            })
    return hits


def check_quote_solution_consistency(quote_path: str, solution_text: str) -> list:
    """检查报价与方案的一致性"""
    issues = []
    quote_modules = extract_modules_from_xlsx(quote_path)

    # 提取关键模块关键词在方案中查找
    module_keywords = {
        '系统基础平台': '平台',
        '能管云盒': '云盒|网关|边缘',
        '通讯箱': '通讯箱',
        '三相多功能电表': '电表|ADL400',
        '数据点位': '点位',
        '批次单耗分析': '批次|单耗',
        '移动端': '移动|小程序|微信',
    }

    for item in quote_modules:
        mod_name = item['module']
        # 查找对应关键词
        for key, pattern in module_keywords.items():
            if key in mod_name:
                if not re.search(pattern, solution_text):
                    # 只在模块数量>0且报价中确实有此模块时报
                    if item.get('quantity', 0) > 0:
                        issues.append(f'[一致性] 报价中有"{mod_name}"，但方案中未找到对应描述（关键词: {pattern}）')
                break

    return issues


def validate(quote_std_path: str, quote_ch_path: str, solution_path: str,
             bottom_sheet_path: str, output_report_path: str) -> dict:
    """执行完整校验"""

    report_lines = []
    all_issues = []
    summary = {'quote_std': 0, 'quote_ch': 0, 'solution': 0, 'consistency': 0}

    report_lines.append('=' * 60)
    report_lines.append('标准能管报价与方案 — 校验报告')
    report_lines.append('=' * 60)
    report_lines.append('')

    # ---- 1. 检查标准价报价 ----
    report_lines.append('【一、标准价报价禁词检查】')
    report_lines.append(f'文件: {quote_std_path}')
    if Path(quote_std_path).exists():
        text = extract_text_from_xlsx(quote_std_path)
        hits = check_forbidden_words(text, FORBIDDEN_WORDS_QUOTE, '报价_标准价')
        if hits:
            for h in hits:
                report_lines.append(f'  ❌ 发现禁词: "{h["word"]}" → {h["context"]}')
            summary['quote_std'] = len(hits)
            all_issues.extend(hits)
        else:
            report_lines.append('  ✅ 未发现禁词')
    else:
        report_lines.append('  ⚠️ 文件不存在，跳过检查')
    report_lines.append('')

    # ---- 2. 检查渠道价报价 ----
    report_lines.append('【二、渠道价报价禁词检查】')
    report_lines.append(f'文件: {quote_ch_path}')
    if Path(quote_ch_path).exists():
        text = extract_text_from_xlsx(quote_ch_path)
        hits = check_forbidden_words(text, FORBIDDEN_WORDS_QUOTE, '报价_渠道价')
        if hits:
            for h in hits:
                report_lines.append(f'  ❌ 发现禁词: "{h["word"]}" → {h["context"]}')
            summary['quote_ch'] = len(hits)
            all_issues.extend(hits)
        else:
            report_lines.append('  ✅ 未发现禁词')
    else:
        report_lines.append('  ⚠️ 文件不存在，跳过检查')
    report_lines.append('')

    # ---- 3. 检查方案 docx ----
    report_lines.append('【三、方案禁词检查】')
    report_lines.append(f'文件: {solution_path}')
    if Path(solution_path).exists():
        text = extract_text_from_docx(solution_path)
        hits = check_forbidden_words(text, FORBIDDEN_WORDS_SOLUTION, '技术方案')
        if hits:
            for h in hits:
                report_lines.append(f'  ❌ 发现禁词: "{h["word"]}" → {h["context"]}')
            summary['solution'] = len(hits)
            all_issues.extend(hits)
        else:
            report_lines.append('  ✅ 未发现禁词')
    else:
        report_lines.append('  ⚠️ 文件不存在，跳过检查')
    report_lines.append('')

    # ---- 4. 报价-方案一致性检查 ----
    report_lines.append('【四、报价-方案一致性检查】')
    if Path(quote_std_path).exists() and Path(solution_path).exists():
        solution_text = extract_text_from_docx(solution_path)
        issues = check_quote_solution_consistency(quote_std_path, solution_text)
        if issues:
            for issue in issues:
                report_lines.append(f'  ⚠️ {issue}')
            summary['consistency'] = len(issues)
        else:
            report_lines.append('  ✅ 关键模块均已覆盖')
    else:
        report_lines.append('  ⚠️ 报价或方案文件不存在，跳过一致性检查')
    report_lines.append('')

    # ---- 5. 内部底稿状态检查 ----
    report_lines.append('【五、内部底稿状态摘要】')
    if Path(bottom_sheet_path).exists():
        data = json.loads(Path(bottom_sheet_path).read_text(encoding='utf-8'))
        status = data.get('status_summary', {})
        totals = data.get('totals', {})
        report_lines.append(f'  标准价总计: {totals.get("grand_total_standard", "N/A"):,.2f}' if isinstance(totals.get('grand_total_standard'), (int, float)) else f'  标准价总计: {totals.get("grand_total_standard", "N/A")}')
        report_lines.append(f'  渠道价总计: {totals.get("grand_total_channel", "N/A"):,.2f}' if isinstance(totals.get('grand_total_channel'), (int, float)) else f'  渠道价总计: {totals.get("grand_total_channel", "N/A")}')
        report_lines.append(f'  状态: AUTO={status.get("AUTO", 0)}, CONFIRM={status.get("CONFIRM", 0)}, '
                           f'INQUIRY={status.get("INQUIRY", 0)}, NONSTANDARD={status.get("NONSTANDARD", 0)}, '
                           f'BLOCKED={status.get("BLOCKED", 0)}')
        if status.get('CONFIRM', 0) > 0:
            report_lines.append('  ⚠️ 存在待确认项，建议人工审核后再正式外发')
        if status.get('INQUIRY', 0) > 0:
            report_lines.append('  ⚠️ 存在待询价项，不能声称完整正式报价')
        if status.get('NONSTANDARD', 0) > 0:
            report_lines.append('  ⚠️ 存在非标项，已超出标准能管路径')
        if status.get('BLOCKED', 0) > 0:
            report_lines.append('  ❌ 存在阻塞项，报价不能外发')
    else:
        report_lines.append('  ⚠️ 内部底稿文件不存在')
    report_lines.append('')

    # ---- 总结 ----
    report_lines.append('=' * 60)
    total_issues = summary['quote_std'] + summary['quote_ch'] + summary['solution'] + summary['consistency']
    report_lines.append(f'校验完成: 共发现 {total_issues} 个问题')
    if total_issues == 0:
        report_lines.append('状态: ✅ 通过 — 未发现禁词或一致性问题')
    else:
        report_lines.append(f'  - 标准价报价禁词: {summary["quote_std"]}')
        report_lines.append(f'  - 渠道价报价禁词: {summary["quote_ch"]}')
        report_lines.append(f'  - 方案禁词: {summary["solution"]}')
        report_lines.append(f'  - 一致性问题: {summary["consistency"]}')
    report_lines.append('=' * 60)

    # 写入报告
    report_text = '\n'.join(report_lines)
    Path(output_report_path).parent.mkdir(parents=True, exist_ok=True)
    Path(output_report_path).write_text(report_text, encoding='utf-8')

    return {
        'total_issues': total_issues,
        'summary': summary,
    }


if __name__ == '__main__':
    quote_std = sys.argv[1] if len(sys.argv) > 1 else 'outputs/quote/报价_标准价.xlsx'
    quote_ch = sys.argv[2] if len(sys.argv) > 2 else 'outputs/quote/报价_渠道价.xlsx'
    solution = sys.argv[3] if len(sys.argv) > 3 else 'outputs/quote/标准技术方案.docx'
    bottom = sys.argv[4] if len(sys.argv) > 4 else 'outputs/quote/internal_bottom_sheet.json'
    report = sys.argv[5] if len(sys.argv) > 5 else 'outputs/quote/校验报告.txt'

    result = validate(quote_std, quote_ch, solution, bottom, report)
    print(f'[OK] 校验报告: {report}')
    print(f'    总问题数: {result["total_issues"]}')
    for k, v in result['summary'].items():
        print(f'    {k}: {v}')
