#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Dump all cell values in Section 8 area of the survey form.
Writes output to a UTF-8 text file for reliable reading.
"""

import openpyxl
import os
import glob

# Find the survey file
SURVEY_PATH = os.path.join(os.path.expanduser("~"), "Desktop", "\u201c\u62a5\u4ef7\u201d\u8c03\u7814\u8868_V1.4.xlsx")
if not os.path.exists(SURVEY_PATH):
    candidates = glob.glob(os.path.join(os.path.expanduser("~"), "Desktop", "*调研表*V1.4*"))
    if candidates:
        SURVEY_PATH = candidates[0]

OUTPUT_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "outputs", "section8_dump.txt")
os.makedirs(os.path.dirname(OUTPUT_FILE), exist_ok=True)

def main():
    lines = []
    def pr(s=""):
        lines.append(s)

    pr(f"Reading: {SURVEY_PATH}")
    pr(f"File exists: {os.path.exists(SURVEY_PATH)}")

    wb = openpyxl.load_workbook(SURVEY_PATH, data_only=True)
    ws = wb.active
    pr(f"Sheet: {ws.title}")
    pr(f"Max row: {ws.max_row}, Max col: {ws.max_column}")

    col_letters = ['A', 'B', 'C', 'D', 'E', 'F', 'G']

    # ===== SECTION 1: Scan for section 8 markers =====
    pr()
    pr("=" * 80)
    pr("SCAN: ALL ROWS WITH SECTION 8 KEYWORDS")
    pr("=" * 80)
    keywords = ['八', '板块八', '可选', '超出默认', '额外', '大屏', '组态', '报表', '移动端', '批次', '软件']
    for row in range(1, ws.max_row + 1):
        for c in col_letters:
            cell = ws[f"{c}{row}"]
            if cell.value is not None:
                val = str(cell.value).strip()
                if any(kw in val for kw in keywords):
                    pr(f"  Row {row}, Col {c}: {repr(cell.value)}")

    # ===== SECTION 2: Full dump rows 73-103 =====
    pr()
    pr("=" * 80)
    pr("FULL CELL DUMP: ROWS 73-103, COLUMNS A-G")
    pr("=" * 80)
    for row in range(73, min(104, ws.max_row + 1)):
        cells = []
        any_val = False
        for c in col_letters:
            cell = ws[f"{c}{row}"]
            if cell.value is not None:
                val = str(cell.value).strip()
                if val:
                    any_val = True
                    cells.append(f"{c}={repr(cell.value)}")
                else:
                    cells.append(f"{c}=''")
            else:
                cells.append(f"{c}=None")
        # Always print rows in the range, even if empty
        pr(f"  Row {row}: {', '.join(cells)}")

    # ===== SECTION 3: Merged cells =====
    pr()
    pr("=" * 80)
    pr("ALL MERGED CELLS IN SHEET (sorted)")
    pr("=" * 80)
    for merged_range in sorted(ws.merged_cells.ranges, key=lambda r: (r.min_row, r.min_col)):
        val = ws.cell(merged_range.min_row, merged_range.min_col).value
        if val is not None:
            val_repr = repr(val)[:80]
        else:
            val_repr = "None"
        pr(f"  {merged_range} = {val_repr}")

    # ===== SECTION 4: Complete dump of ENTIRE sheet (all non-empty rows) =====
    pr()
    pr("=" * 80)
    pr("ALL NON-EMPTY ROWS IN SHEET")
    pr("=" * 80)
    for row in range(1, ws.max_row + 1):
        cells = []
        any_val = False
        for c in col_letters:
            cell = ws[f"{c}{row}"]
            if cell.value is not None:
                val = str(cell.value).strip()
                if val:
                    any_val = True
                    cells.append(f"{c}={repr(cell.value)}")
        if any_val:
            pr(f"  Row {row}: {', '.join(cells)}")

    wb.close()
    pr()
    pr("Done!")

    # Write output
    with open(OUTPUT_FILE, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))

    print(f"Output written to: {OUTPUT_FILE}")

if __name__ == "__main__":
    main()
