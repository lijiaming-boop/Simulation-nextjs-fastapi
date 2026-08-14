"""业务服务层：封装 core 模块，供 FastAPI 路由调用"""
import asyncio
import json
import os
import subprocess
import sys
import threading
from datetime import datetime
from pathlib import Path
from typing import Optional

from core.skill_caller import SkillCaller, SkillError
from core.data_filler import DataFiller
from core.excel_calculator import ExcelCalculator

# 模块加载时缓存 API Key（供后台线程创建 SkillCaller 时显式传入）
_LLM_API_KEY = os.getenv("ANTHROPIC_API_KEY", "")
from config import WeatherService

from .deps import PROJECT_ROOT, session_store


# ==================== 任务注册表（后台线程） ====================

# session_id → 后台 thread 对象
_extract_tasks: dict = {}
# session_id → calculate 后台 thread
_calculate_tasks: dict = {}
# session_id → suggest 后台 thread
_suggest_tasks: dict = {}
# session_id → 仿真任务后台 thread
_sim_tasks: dict = {}
# session_id → 报价提取后台 thread
_quote_extract_tasks: dict = {}
# session_id → 报价生成后台 thread
_quote_generate_tasks: dict = {}
# session_id → 报价计算后台 thread
_quote_calculate_tasks: dict = {}


def is_extract_running(session_id: str) -> bool:
    """查询提取任务是否仍在运行"""
    t = _extract_tasks.get(session_id)
    return t is not None and t.is_alive()


def is_calculate_running(session_id: str) -> bool:
    t = _calculate_tasks.get(session_id)
    return t is not None and t.is_alive()


def is_suggest_running(session_id: str) -> bool:
    t = _suggest_tasks.get(session_id)
    return t is not None and t.is_alive()


def is_sim_running(session_id: str) -> bool:
    t = _sim_tasks.get(session_id)
    return t is not None and t.is_alive()


def is_quote_extract_running(session_id: str) -> bool:
    t = _quote_extract_tasks.get(session_id)
    return t is not None and t.is_alive()


def is_quote_generate_running(session_id: str) -> bool:
    t = _quote_generate_tasks.get(session_id)
    return t is not None and t.is_alive()


def is_quote_calculate_running(session_id: str) -> bool:
    t = _quote_calculate_tasks.get(session_id)
    return t is not None and t.is_alive()


# ==================== AI 事件通道 ====================

def ensure_ai_queue(session_id: str) -> asyncio.Queue:
    """为 session 创建/复用 AI 事件 Queue"""
    s = session_store.get(session_id)
    if s is None:
        raise RuntimeError(f"session {session_id!r} 不存在")
    if "ai_queue" not in s:
        # 注意：必须在主线程事件循环中创建，才能被 SSE 端点消费
        s["ai_queue"] = asyncio.Queue()
    return s["ai_queue"]


def get_ai_queue(session_id: str) -> Optional[asyncio.Queue]:
    s = session_store.get(session_id)
    return s.get("ai_queue") if s else None


# ==================== 提取任务（在线程中执行） ====================

def _run_extract_in_thread(
    session_id: str,
    module_type: str,
    input_path: Path,
    output_json_path: Path,
    original_filename: str,
) -> None:
    """在线程中调用 SkillCaller 完成数据提取

    所有 AI 消息通过 asyncio.Queue 推送给 SSE 端点
    """
    # 拿到主线程事件循环（用于 put_nowait）
    loop = session_store._main_loop  # 由 extract() 端点注入

    def ai_callback(timestamp: str, message: str) -> None:
        """SkillCaller 回调：推事件到 AI Queue + 累积到 ai_logs"""
        session_store.append_ai_log(session_id, timestamp, message)
        event = {"type": "log", "timestamp": timestamp, "message": message}
        q = get_ai_queue(session_id)
        if q is not None:
            loop.call_soon_threadsafe(q.put_nowait, event)

    try:
        # === 空调场景：Claude Agent SDK 调技能（脚本为纯 Python，无需 LLM） ===
        if module_type == "air_conditioner":
            ai_callback(datetime.now().strftime("%H:%M:%S"), "🚀 开始提取数据（Claude Agent SDK）...")

            caller = SkillCaller(timeout=600, api_key=_LLM_API_KEY or None, project_root=PROJECT_ROOT)
            caller.call_skill(
                "extract-and-check-data",
                str(input_path),
                str(output_json_path),
                original_filename,
                callback=ai_callback,
            )

            if not output_json_path.exists():
                raise SkillError(f"技能执行完成但未生成输出: {output_json_path}")
        else:
            # === 空压场景：Claude Agent SDK 一次完成文本提取 + 结构化 ===
            ai_callback(datetime.now().strftime("%H:%M:%S"), "🚀 [空压] 提取结构化数据（Claude Agent SDK）...")

            caller = SkillCaller(timeout=600, api_key=_LLM_API_KEY or None, project_root=PROJECT_ROOT)
            caller.call_skill(
                "extract-air-compressor-data-llm",
                str(input_path),
                str(output_json_path),
                original_filename,
                callback=ai_callback,
            )

            if not output_json_path.exists():
                raise SkillError(f"技能执行完成但未生成输出: {output_json_path}")

        # 读取输出 JSON
        with open(output_json_path, "r", encoding="utf-8") as f:
            content = json.load(f)
        s = session_store.get(session_id)
        if s is not None:
            # 兼容两种 schema：带 extracted_data/quality_report 的拆分 或 整体
            if "extracted_data" in content and "quality_report" in content:
                s["extracted_data"] = content["extracted_data"]
                s["quality_report"] = content["quality_report"]
            elif "rawData" in content:  # 空压 LLM 版
                s["extracted_data"] = content.get("rawData")
                s["quality_report"] = content.get("quality_report") or content.get("rejectionCheck")
            else:
                # 兜底：整体作为 extracted_data
                s["extracted_data"] = content

        end_ts = datetime.now().strftime("%H:%M:%S")
        ai_callback(end_ts, f"✅ 提取完成，输出: {output_json_path.name}")
        session_store.update_stage(session_id, "extract_check")

        # 发 done 事件
        if (q := get_ai_queue(session_id)) is not None:
            loop.call_soon_threadsafe(
                q.put_nowait,
                {"type": "done", "timestamp": datetime.now().strftime("%H:%M:%S")},
            )

    except Exception as e:
        err_ts = datetime.now().strftime("%H:%M:%S")
        err_msg = f"❌ {type(e).__name__}: {e}"
        session_store.append_ai_log(session_id, err_ts, err_msg)
        if (q := get_ai_queue(session_id)) is not None:
            loop.call_soon_threadsafe(
                q.put_nowait,
                {"type": "error", "timestamp": err_ts, "message": err_msg},
            )
        session_store.update_stage(session_id, "error")


# ==================== 提取入口 ====================

def start_extract(
    session_id: str,
    module_type: str,
    input_path: Path,
    original_filename: str,
    main_loop: asyncio.AbstractEventLoop,
) -> Path:
    """启动后台提取任务，返回输出 JSON 路径"""
    if is_extract_running(session_id):
        raise RuntimeError(f"session {session_id!r} 已有提取任务在运行")

    # 输出 JSON 命名：沿用 Streamlit 惯例
    if module_type == "air_conditioner":
        output_json_path = session_store.get(session_id)["session_dir"] / "extracted_data.json"
    else:
        output_json_path = session_store.get(session_id)["session_dir"] / "extracted_air_compressor_data.json"

    # 注入主事件循环（线程内 put_nowait 需要用）
    session_store._main_loop = main_loop

    # 确保 AI Queue 存在
    ensure_ai_queue(session_id)

    session_store.update_stage(session_id, "extracting")

    t = threading.Thread(
        target=_run_extract_in_thread,
        args=(session_id, module_type, input_path, output_json_path, original_filename),
        daemon=True,
        name=f"extract-{session_id}",
    )
    t.start()
    _extract_tasks[session_id] = t

    return output_json_path


# ==================== 计算任务 ====================

def _run_calculate_in_thread(
    session_id: str,
    module_type: str,
    extracted_data: dict,
    filling_suggestions: Optional[dict],
) -> None:
    """后台执行 DataFiller + ExcelCalculator"""
    loop = session_store._main_loop

    def emit(event_type: str, message: str) -> None:
        ts = datetime.now().strftime("%H:%M:%S")
        session_store.append_ai_log(session_id, ts, message)
        q = get_ai_queue(session_id)
        if q is not None:
            loop.call_soon_threadsafe(
                q.put_nowait,
                {"type": event_type, "timestamp": ts, "message": message},
            )

    try:
        s = session_store.get(session_id)
        if s is None:
            raise RuntimeError(f"session {session_id!r} 不存在")
        session_dir: Path = s["session_dir"]

        if module_type == "air_compressor":
            # ==================== 空压机计算路径 ====================
            from core.air_compressor_calculator import AirCompressorCalculator

            # 1. 将 LLM 提取结构映射为 Calculator 期望的格式
            emit("log", "📄 映射空压机数据结构...")
            project_level = extracted_data.get("projectLevel", {}) if isinstance(extracted_data, dict) else {}
            raw_equipment = extracted_data.get("equipmentLevel", []) if isinstance(extracted_data, dict) else []

            # 设备类型名映射：螺杆→螺杆式、离心→离心式
            type_map = {"螺杆": "螺杆式", "离心": "离心式"}
            equipment_list = []
            for eq in (raw_equipment if isinstance(raw_equipment, list) else []):
                mapped = dict(eq)
                et = mapped.get("equipmentType", "")
                if et in type_map:
                    mapped["equipmentType"] = type_map[et]
                equipment_list.append(mapped)

            calculator_input = {
                "system_info": {
                    "project_name": project_level.get("projectName", ""),
                    "customer_name": project_level.get("companyName", ""),
                    "contact_person": "",
                    "contact_phone": "",
                },
                "equipment_list": equipment_list,
                "operation_data": {
                    "annual_operating_hours": (
                        (project_level.get("dailyRunningHours") or 10)
                        * (project_level.get("annualRunningDays") or 330)
                    ),
                },
                "completeness": 1.0,
                "filled_count": 0,
                "session_dir": str(session_dir),
                "original_filename": s.get("original_filename", ""),
            }

            filled_path = session_dir / "compressor_filled_data.json"
            with open(filled_path, "w", encoding="utf-8") as f:
                json.dump(calculator_input, f, ensure_ascii=False, indent=2)

            # 2. 应用填充建议（如果有）
            if filling_suggestions and filling_suggestions.get("fillable_suggestions"):
                emit("log", "📝 应用空压填充建议...")
                for sug in filling_suggestions.get("fillable_suggestions", []):
                    field = sug.get("field", "")
                    value = sug.get("suggested_value")
                    if value is not None and field and not field.startswith("equipmentList["):
                        calculator_input[field] = value
                with open(filled_path, "w", encoding="utf-8") as f:
                    json.dump(calculator_input, f, ensure_ascii=False, indent=2)

            # 3. AirCompressorCalculator 计算
            emit("log", "⚙️ 运行空压机节能计算...")
            calculator = AirCompressorCalculator()
            result_path = str(session_dir / "compressor_result.json")
            calculator.calculate(str(filled_path), result_path)
            emit("log", "✅ 计算完成")

            # 4. 读取结果
            with open(result_path, "r", encoding="utf-8") as f:
                result = json.load(f)
            with open(filled_path, "r", encoding="utf-8") as f:
                filled_data = json.load(f)

            formatted = calculator.format_export_results(filled_data, result)

            # 5. 写回 session_store
            s["filling_suggestions"] = filling_suggestions
            s["result"] = result
            s["formatted_result"] = formatted
            session_store.update_stage(session_id, "result")

        else:
            # ==================== 空调计算路径（原有逻辑） ====================

            # 1. 写入提取数据文件
            emit("log", "📄 写入提取数据...")
            extracted_path = session_dir / "ac_extracted_data.json"
            with open(extracted_path, "w", encoding="utf-8") as f:
                json.dump(extracted_data, f, ensure_ascii=False, indent=2)

            # 2. 写入填充建议（默认空建议）
            suggestions_path = session_dir / "ac_filling_suggestions.json"
            if filling_suggestions is None:
                filling_suggestions = {
                    "can_proceed": True,
                    "fillable_suggestions": [],
                    "unfillable_fields": [],
                }
            with open(suggestions_path, "w", encoding="utf-8") as f:
                json.dump(filling_suggestions, f, ensure_ascii=False, indent=2)

            # 3. DataFiller 执行填充
            emit("log", "🔄 执行数据填充...")
            filler = DataFiller()
            filled_path = filler.fill_data(
                str(extracted_path),
                str(suggestions_path),
                str(session_dir / "ac_filled_data.json"),
            )
            emit("log", "✅ 填充完成")

            # 4. ExcelCalculator 执行公式计算
            emit("log", "⚙️ 运行 Excel 公式计算引擎...")
            weather = WeatherService()
            calculator = ExcelCalculator(weather_service=weather)
            result_path = calculator.calculate(
                filled_path,
                str(session_dir / "ac_result_data.json"),
            )
            emit("log", "✅ 计算完成")

            # 5. 格式化输出
            with open(filled_path, "r", encoding="utf-8") as f:
                filled_data = json.load(f)
            with open(result_path, "r", encoding="utf-8") as f:
                result = json.load(f)

            formatted = calculator.format_export_results(filled_data, result)
            formatted_path = session_dir / "ac_formatted_result.json"
            with open(formatted_path, "w", encoding="utf-8") as f:
                json.dump(formatted, f, ensure_ascii=False, indent=2)

            # 5.5 生成报价单和方案文档
            emit("log", "生成报价单和方案文档...")
            calculator._generate_quotation(str(session_dir), filled_data, result)
            calculator._generate_scheme_document(str(session_dir), filled_data, result)
            emit("log", "报价单生成完成")

            # 6. 写回 session_store
            s["filling_suggestions"] = filling_suggestions
            s["result"] = result
            s["formatted_result"] = formatted
            session_store.update_stage(session_id, "result")

        # 7. 完成事件
        q = get_ai_queue(session_id)
        if q is not None:
            loop.call_soon_threadsafe(
                q.put_nowait,
                {"type": "done", "timestamp": datetime.now().strftime("%H:%M:%S")},
            )

    except Exception as e:
        err_ts = datetime.now().strftime("%H:%M:%S")
        err_msg = f"❌ {type(e).__name__}: {e}"
        session_store.append_ai_log(session_id, err_ts, err_msg)
        if (q := get_ai_queue(session_id)) is not None:
            loop.call_soon_threadsafe(
                q.put_nowait,
                {"type": "error", "timestamp": err_ts, "message": err_msg},
            )
        # 将错误信息存入 session，供 /api/result 返回
        s_err = session_store.get(session_id)
        if s_err is not None:
            s_err["calc_error"] = f"{type(e).__name__}: {e}"
        session_store.update_stage(session_id, "error")


def start_calculate(
    session_id: str,
    module_type: str,
    extracted_data: dict,
    filling_suggestions: Optional[dict],
    main_loop: asyncio.AbstractEventLoop,
) -> None:
    """启动后台计算任务"""
    if is_calculate_running(session_id):
        raise RuntimeError(f"session {session_id!r} 已有计算任务在运行")

    session_store._main_loop = main_loop
    ensure_ai_queue(session_id)
    session_store.update_stage(session_id, "calculating")

    t = threading.Thread(
        target=_run_calculate_in_thread,
        args=(session_id, module_type, extracted_data, filling_suggestions),
        daemon=True,
        name=f"calculate-{session_id}",
    )
    t.start()
    _calculate_tasks[session_id] = t


# ==================== 填充建议任务 ====================


def _run_suggest_in_thread(
    session_id: str,
    module_type: str,
    use_llm: bool,
) -> None:
    """后台生成数据填充建议

    1. 把内存里的 extracted_data / quality_report 写盘
    2. subprocess 调脚本生成规则建议（区域电价 + 季节默认值）
    3. use_llm=True 时，再用 SkillCaller(SDK) 对规则建议做智能增强
    """
    loop = session_store._main_loop
    s = session_store.get(session_id)
    if s is None:
        return
    session_dir = s["session_dir"]

    def emit(event_type: str, message: str) -> None:
        ts = datetime.now().strftime("%H:%M:%S")
        session_store.append_ai_log(session_id, ts, message)
        q = get_ai_queue(session_id)
        if q is not None:
            loop.call_soon_threadsafe(
                q.put_nowait,
                {"type": event_type, "timestamp": ts, "message": message},
            )

    try:
        extracted_data = s.get("extracted_data")
        quality_report = s.get("quality_report") or {"completeness": 1.0, "issues": []}
        if extracted_data is None:
            raise RuntimeError("extracted_data 缺失，请先调用 /extract")

        # 1. 写盘
        emit("log", "📄 写入提取数据 + 质量报告...")
        prefix = "compressor_" if module_type == "air_compressor" else "ac_"
        extracted_path = session_dir / f"{prefix}extracted_data.json"
        with open(extracted_path, "w", encoding="utf-8") as f:
            json.dump(extracted_data, f, ensure_ascii=False, indent=2)
        quality_path = session_dir / f"{prefix}quality_report.json"
        with open(quality_path, "w", encoding="utf-8") as f:
            json.dump(quality_report, f, ensure_ascii=False, indent=2)
        suggestions_path = session_dir / f"{prefix}filling_suggestions.json"

        # 2. 规则脚本（按模块类型分发）
        if module_type == "air_compressor":
            skill_dir = "suggest-air-compressor-data-filling"
            script_name = "suggest_air_compressor_data_filling.py"
            emit("log", "🔧 调用空压规则脚本 suggest_air_compressor_data_filling.py ...")
        else:
            skill_dir = "suggest-data-filling"
            script_name = "suggest_data_filling.py"
            emit("log", "🔧 调用空调规则脚本 suggest_data_filling.py ...")
        script = PROJECT_ROOT / ".claude" / "skills" / skill_dir / "scripts" / script_name
        cmd = [
            sys.executable,
            str(script),
            str(extracted_path),
            str(quality_path),
            str(suggestions_path),
        ]
        proc = subprocess.Popen(
            cmd,
            cwd=str(PROJECT_ROOT),
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            encoding="utf-8",
            errors="replace",
            env={**os.environ, "PYTHONIOENCODING": "utf-8", "PYTHONUTF8": "1"},
        )
        for line in proc.stdout:
            line = line.rstrip("\n")
            if line:
                emit("log", line)
        proc.wait()
        if proc.returncode != 0:
            stderr = proc.stderr.read() if proc.stderr else ""
            raise SkillError(
                f"规则脚本退出码 {proc.returncode}: {stderr[:500]}"
            )
        if not suggestions_path.exists():
            raise SkillError("规则脚本执行完成但未生成输出")

        with open(suggestions_path, "r", encoding="utf-8") as f:
            suggestions = json.load(f)
        emit("log", "✅ 规则建议生成完成")

        # 3. LLM 增强（可选，走 Claude Agent SDK）
        if use_llm:
            emit("log", "🤖 调用 LLM 对规则建议做智能增强...")
            try:
                client = SkillCaller(timeout=180, api_key=_LLM_API_KEY or None, project_root=PROJECT_ROOT)
                if module_type == "air_compressor":
                    prompt = (
                        "你是一位空压机节能测算顾问。下面是某空压机项目的提取数据、质量报告以及基于规则生成的填充建议。\n"
                        f"【提取数据】\n{json.dumps(extracted_data, ensure_ascii=False)[:1800]}\n\n"
                        f"【质量报告】\n{json.dumps(quality_report, ensure_ascii=False)[:600]}\n\n"
                        f"【规则建议】\n{json.dumps(suggestions, ensure_ascii=False)[:1800]}\n\n"
                        "请给出 2-4 条更智能的填充建议（例如行业默认压力范围、设备效率曲线、联控策略），"
                        "每条建议一行，限 200 字以内。"
                    )
                else:
                    prompt = (
                        "你是一位暖通节能测算顾问。下面是某中央空调项目的提取数据、质量报告以及基于规则生成的填充建议。\n"
                        f"【提取数据】\n{json.dumps(extracted_data, ensure_ascii=False)[:1800]}\n\n"
                        f"【质量报告】\n{json.dumps(quality_report, ensure_ascii=False)[:600]}\n\n"
                        f"【规则建议】\n{json.dumps(suggestions, ensure_ascii=False)[:1800]}\n\n"
                        "请给出 2-4 条更智能的填充建议（例如季节插值策略、行业默认值、电价区间），"
                        "每条建议一行，限 200 字以内。"
                    )
                llm_text = client.create_message(
                    prompt,
                    system="你是节能测算顾问，回复简洁专业。",
                    max_tokens=400,
                )
                suggestions["llm_enhanced_tips"] = llm_text
                emit("log", f"✅ LLM 增强完成：{llm_text[:80]}...")
            except Exception as llm_err:
                emit(
                    "log",
                    f"⚠️ LLM 增强失败（保留规则建议）：{type(llm_err).__name__}: {llm_err}",
                )

        # 4. 写回 session_store
        s["filling_suggestions"] = suggestions
        with open(suggestions_path, "w", encoding="utf-8") as f:
            json.dump(suggestions, f, ensure_ascii=False, indent=2)
        session_store.update_stage(session_id, "suggest")

        # 5. done 事件
        q = get_ai_queue(session_id)
        if q is not None:
            loop.call_soon_threadsafe(
                q.put_nowait,
                {"type": "done", "timestamp": datetime.now().strftime("%H:%M:%S")},
            )

    except Exception as e:
        err_ts = datetime.now().strftime("%H:%M:%S")
        err_msg = f"❌ {type(e).__name__}: {e}"
        session_store.append_ai_log(session_id, err_ts, err_msg)
        if (q := get_ai_queue(session_id)) is not None:
            loop.call_soon_threadsafe(
                q.put_nowait,
                {"type": "error", "timestamp": err_ts, "message": err_msg},
            )
        session_store.update_stage(session_id, "error")


def start_suggest(
    session_id: str,
    module_type: str,
    use_llm: bool,
    main_loop: asyncio.AbstractEventLoop,
) -> None:
    """启动后台填充建议任务"""
    if is_suggest_running(session_id):
        raise RuntimeError(f"session {session_id!r} 已有建议任务在运行")

    session_store._main_loop = main_loop
    ensure_ai_queue(session_id)
    session_store.update_stage(session_id, "suggesting")

    t = threading.Thread(
        target=_run_suggest_in_thread,
        args=(session_id, module_type, use_llm),
        daemon=True,
        name=f"suggest-{session_id}",
    )
    t.start()
    _suggest_tasks[session_id] = t


# ==================== 仿真模块服务函数 ====================


def _start_sim_task(
    session_id: str,
    task_func,
    main_loop: asyncio.AbstractEventLoop,
    **kwargs,
) -> threading.Thread:
    """启动仿真后台任务线程"""
    session_store._main_loop = main_loop
    t = threading.Thread(target=task_func, kwargs={"session_id": session_id, **kwargs}, daemon=True)
    _sim_tasks[session_id] = t
    t.start()
    return t


def _sim_callback(session_id: str, loop: asyncio.AbstractEventLoop, timestamp: str, message: str) -> None:
    """仿真任务 AI 回调"""
    session_store.append_ai_log(session_id, timestamp, message)
    event = {"type": "log", "timestamp": timestamp, "message": message}
    q = get_ai_queue(session_id)
    if q is not None:
        loop.call_soon_threadsafe(q.put_nowait, event)


def _sim_done(session_id: str, loop: asyncio.AbstractEventLoop, error: str = None) -> None:
    """仿真任务完成事件"""
    event = {
        "type": "error" if error else "done",
        "timestamp": datetime.now().strftime("%H:%M:%S"),
    }
    if error:
        event["message"] = error
    q = get_ai_queue(session_id)
    if q is not None:
        loop.call_soon_threadsafe(q.put_nowait, event)


# --- 拓扑提取 ---

def _extract_project_name_from_survey(excel_path) -> str:
    """从调研表 Excel 中提取项目名称（纯 Python，不依赖 AI）"""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(excel_path), read_only=True, data_only=True)
    except Exception:
        return ""
    labels = ["项目名称", "站房名称", "工程名称", "项目名", "站房名",
              "Project", "project", "Site", "site"]
    try:
        for sheet_name in list(wb.sheetnames[:3]):
            ws = wb[sheet_name]
            for row in ws.iter_rows(min_row=1, max_row=20, max_col=15, values_only=False):
                cells = list(row)
                for i, cell in enumerate(cells):
                    val = cell.value
                    if val and isinstance(val, str):
                        val_stripped = val.strip()
                        for label in labels:
                            if val_stripped == label:
                                if i + 1 < len(cells):
                                    right_val = cells[i + 1].value
                                    if right_val and isinstance(right_val, str) and right_val.strip():
                                        wb.close()
                                        return right_val.strip()
                        for label in labels:
                            if label in val_stripped and len(val_stripped) > len(label):
                                for sep in ["：", ":", "="]:
                                    if sep in val_stripped:
                                        name = val_stripped.split(sep, 1)[1].strip()
                                        if name and name not in labels:
                                            wb.close()
                                            return name
        wb.close()
    except Exception:
        pass
    return ""


def _run_topology_extraction(session_id: str, upload_path: str, loop: asyncio.AbstractEventLoop) -> None:
    """后台线程：从调研表提取拓扑"""
    s = session_store.get(session_id)
    session_dir = s["session_dir"]
    output_path = session_dir / "sim_topology.json"
    # 保存用户已编辑的项目名称（AI 提取不应覆盖用户输入）
    user_project_name = (s.get("sim_project_name") or "").strip()

    def cb(ts, msg):
        _sim_callback(session_id, loop, ts, msg)

    try:
        # 先将 Excel 全部内容读为文本，确保 AI 能看到完整数据
        excel_text_path = session_dir / "survey_text_dump.txt"
        try:
            import pandas as pd
            cb(datetime.now().strftime("%H:%M:%S"), "📄 预读调研表内容...")
            xls = pd.ExcelFile(upload_path)
            lines = []
            for sheet_name in xls.sheet_names:
                df = pd.read_excel(xls, sheet_name=sheet_name, header=None)
                lines.append(f"\n=== Sheet: {sheet_name} ({df.shape[0]} rows x {df.shape[1]} cols) ===")
                lines.append(df.to_string(index=False, na_rep=""))
            excel_text_path.write_text("\n".join(lines), encoding="utf-8")
            cb(datetime.now().strftime("%H:%M:%S"), f"📄 已读取 {len(xls.sheet_names)} 个 sheet")
        except Exception as e:
            cb(datetime.now().strftime("%H:%M:%S"), f"⚠️ 预读Excel失败: {e}，将由AI直接读取")
            excel_text_path = None

        caller = SkillCaller(timeout=300)
        cb(datetime.now().strftime("%H:%M:%S"), "🚀 开始从调研表提取设备拓扑...")

        args = [str(upload_path), str(output_path), "extract_topology"]
        if excel_text_path:
            args.append(str(excel_text_path))

        skill_result = caller.call_skill(
            "plant-simulation",
            *args,
            callback=cb,
            max_retries=2,
        )

        if output_path.exists():
            with open(output_path, "r", encoding="utf-8") as f:
                raw = json.load(f)
            topology = _normalize_topology(raw)

            # 项目名称优先级: 用户已编辑 > AI提取 > 调研表扫描 > 默认
            placeholder_names = {'', '未知项目', '项目名称', 'XXXX', 'xxx', 'xxxx',
                                 'project_name', 'Project Name', 'XX工厂'}
            ai_name = (topology.get("project_name") or "").strip()
            user_is_placeholder = user_project_name in placeholder_names or len(user_project_name) <= 2
            ai_is_placeholder = ai_name in placeholder_names or len(ai_name) <= 2

            if not user_is_placeholder:
                topology["project_name"] = user_project_name
            elif ai_is_placeholder:
                survey_name = _extract_project_name_from_survey(upload_path)
                if survey_name and survey_name not in placeholder_names:
                    topology["project_name"] = survey_name
                else:
                    topology["project_name"] = "未知项目"

            s["sim_topology"] = topology
            s["sim_project_name"] = topology["project_name"]
            s["sim_stage"] = "cop_curves"
            cb(datetime.now().strftime("%H:%M:%S"), f"✅ 拓扑识别完成：{len(topology.get('chillers',[]))} 台主机")
            _sim_done(session_id, loop)
        else:
            _sim_done(session_id, loop, "拓扑识别完成但未生成输出文件")
    except Exception as e:
        _sim_done(session_id, loop, f"拓扑提取失败: {e}")


def _run_topology_parse(session_id: str, text_input: str, project_name: str, loop: asyncio.AbstractEventLoop) -> None:
    """后台线程：文本解析拓扑"""
    s = session_store.get(session_id)
    session_dir = s["session_dir"]
    input_path = session_dir / "sim_topology_input.txt"
    output_path = session_dir / "sim_topology.json"

    def cb(ts, msg):
        _sim_callback(session_id, loop, ts, msg)

    try:
        input_path.write_text(text_input, encoding="utf-8")
        caller = SkillCaller(timeout=180)
        cb(datetime.now().strftime("%H:%M:%S"), "🚀 开始解析拓扑文本...")

        caller.call_skill(
            "plant-simulation",
            str(input_path),
            str(output_path),
            "parse_topology",
            callback=cb,
            max_retries=2,
        )

        if output_path.exists():
            with open(output_path, "r", encoding="utf-8") as f:
                raw = json.load(f)
            topology = _normalize_topology(raw)
            s["sim_topology"] = topology
            s["sim_project_name"] = topology.get("project_name", project_name or "未知项目")
            s["sim_stage"] = "cop_curves"
            cb(datetime.now().strftime("%H:%M:%S"), "✅ 拓扑解析完成")
            _sim_done(session_id, loop)
        else:
            _sim_done(session_id, loop, "拓扑解析完成但未生成输出")
    except Exception as e:
        _sim_done(session_id, loop, f"拓扑解析失败: {e}")


# --- COP 曲线解析 ---

def _run_cop_parse(session_id: str, img_path: str, idx: int, model_name: str, rated_cop: float, loop: asyncio.AbstractEventLoop) -> None:
    """后台线程：图片识别 COP 曲线"""
    s = session_store.get(session_id)
    session_dir = s["session_dir"]
    output_path = session_dir / f"sim_cop_{idx}.json"

    def cb(ts, msg):
        _sim_callback(session_id, loop, ts, msg)

    try:
        caller = SkillCaller(timeout=300)
        cb(datetime.now().strftime("%H:%M:%S"), f"🚀 开始识别 {model_name} COP 曲线...")

        caller.call_skill(
            "plant-simulation",
            str(img_path),
            str(output_path),
            "parse_cop_curve",
            callback=cb,
        )

        if output_path.exists():
            with open(output_path, "r", encoding="utf-8") as f:
                cop_data = json.load(f)
            cw_temps = cop_data.get("cw_temps", [])
            load_ratios = cop_data.get("load_ratios", [])
            cop_table = cop_data.get("cop_table", [])
            if not cw_temps or not load_ratios or not cop_table:
                # fallback to default
                default_cop = _generate_default_cop_table(rated_cop)
                default_cop["_source"] = "auto_fallback"
                s["sim_cop_data"][model_name] = default_cop
                cb(datetime.now().strftime("%H:%M:%S"), "ℹ️ 图片识别未返回有效数据，已使用默认曲线")
            else:
                s["sim_cop_data"][model_name] = cop_data
                cb(datetime.now().strftime("%H:%M:%S"), f"✅ COP 数据已收集 ({len(cw_temps)}×{len(load_ratios)})")
            s["sim_stage"] = "cop_curves"
            _sim_done(session_id, loop)
        else:
            _sim_done(session_id, loop, "COP 识别完成但未生成输出")
    except Exception as e:
        _sim_done(session_id, loop, f"COP 识别失败: {e}")


# --- 工况生成 ---

def _run_operating_generate(session_id: str, text_input: str, loop: asyncio.AbstractEventLoop) -> None:
    """后台线程：生成月度运行工况"""
    s = session_store.get(session_id)
    session_dir = s["session_dir"]
    input_path = session_dir / "sim_operating_input.txt"
    output_path = session_dir / "sim_operating.json"

    def cb(ts, msg):
        _sim_callback(session_id, loop, ts, msg)

    try:
        input_path.write_text(text_input, encoding="utf-8")
        caller = SkillCaller(timeout=300)
        cb(datetime.now().strftime("%H:%M:%S"), "🚀 开始生成月度运行工况...")

        caller.call_skill(
            "plant-simulation",
            str(input_path),
            str(output_path),
            "generate_operating",
            callback=cb,
        )

        if output_path.exists():
            with open(output_path, "r", encoding="utf-8") as f:
                operating = json.load(f)
            s["sim_operating_conditions"] = operating
            s["sim_stage"] = "operating"
            cb(datetime.now().strftime("%H:%M:%S"), "✅ 工况数据已生成")
            _sim_done(session_id, loop)
        else:
            _sim_done(session_id, loop, "工况生成完成但未生成输出")
    except Exception as e:
        _sim_done(session_id, loop, f"工况生成失败: {e}")


# --- 仿真执行 ---

def _run_simulation(session_id: str, loop: asyncio.AbstractEventLoop) -> None:
    """后台线程：执行仿真"""
    s = session_store.get(session_id)
    session_dir = s["session_dir"]
    sim_mode = s["sim_mode"]

    def cb(ts, msg):
        _sim_callback(session_id, loop, ts, msg)

    try:
        # 打包输入
        sim_input = {
            "project_name": s.get("sim_project_name", ""),
            "sim_mode": sim_mode,
            "topology": s.get("sim_topology"),
            "cop_data": s.get("sim_cop_data", {}),
            "operating_conditions": s.get("sim_operating_conditions"),
        }

        # old_station 多方案对比：追加基线耗电量 + 方案列表
        if sim_mode == "old_station":
            baseline = s.get("sim_baseline_power_kwh")
            schemes = s.get("sim_schemes") or []
            if baseline is None or baseline <= 0:
                raise ValueError("old_station 模式缺少基准工况耗电量（sim_baseline_power_kwh）")
            if not schemes:
                raise ValueError("old_station 模式缺少优化方案（sim_schemes 不能为空）")
            sim_input["baseline_power_kwh"] = baseline
            sim_input["schemes"] = schemes
            cb(
                datetime.now().strftime("%H:%M:%S"),
                f"📊 多方案对比模式：基准 {baseline/10000:.2f}万 kWh + {len(schemes)} 个优化方案",
            )

        input_path = session_dir / "sim_input_all.json"
        with open(input_path, "w", encoding="utf-8") as f:
            json.dump(sim_input, f, ensure_ascii=False, indent=2)

        output_dir = session_dir / "sim_output"
        output_dir.mkdir(exist_ok=True)

        caller = SkillCaller(timeout=900)
        cb(datetime.now().strftime("%H:%M:%S"), f"🚀 开始仿真（{sim_mode}）...")

        try:
            caller.call_skill(
                "plant-simulation",
                str(input_path),
                str(output_dir),
                sim_mode,
                callback=cb,
                max_retries=2,
            )
        except Exception as timeout_e:
            # 超时容错：检查输出文件是否已生成
            csv_files = list(output_dir.glob("*.csv"))
            xlsx_files = list(output_dir.glob("*.xlsx"))
            if csv_files or xlsx_files:
                cb(datetime.now().strftime("%H:%M:%S"), "⚠️ 仿真执行超时，但已检测到输出文件，尝试读取结果...")
            else:
                raise

        # 收集结果
        result = {"success": True, "output_dir": str(output_dir), "files": []}
        for f in output_dir.iterdir():
            result["files"].append(f.name)

        report_files = list(output_dir.glob("*报告*.md")) + list(output_dir.glob("*report*.md"))
        report_text = ""
        if report_files:
            report_text = report_files[0].read_text(encoding="utf-8")
            result["report"] = report_text

        summary_files = list(output_dir.glob("*汇总*.json")) + list(output_dir.glob("*summary*.json"))
        if summary_files:
            with open(summary_files[0], "r", encoding="utf-8") as f:
                result["summary"] = json.load(f)
        elif report_text:
            # summary.json 缺失时从报告 Markdown 解析指标
            import re
            summary = {}
            m = re.search(r"基准年耗电量[：:]\s*([\d,.]+)\s*(万)?kWh", report_text)
            if m:
                val = float(m.group(1).replace(',', ''))
                if m.group(2):
                    val *= 10000
                summary["baseline_power_kwh_simulated"] = val
            m = re.search(r"优化年耗电量[：:]\s*([\d,.]+)\s*(万)?kWh", report_text)
            if m:
                val = float(m.group(1).replace(',', ''))
                if m.group(2):
                    val *= 10000
                summary["optimized_power_kwh_simulated"] = val
            m = re.search(r"年节能率[：:]\s*([\d.]+)\s*%", report_text)
            if m:
                summary["actual_energy_saving_rate"] = float(m.group(1)) / 100
            m = re.search(r"年节电量[：:]\s*([\d,.]+)\s*(万)?kWh", report_text)
            if m:
                val = float(m.group(1).replace(',', ''))
                if m.group(2):
                    val *= 10000
                summary["energy_saving_kwh"] = val
            if summary:
                result["summary"] = summary

        # 多方案对比汇总（优先）
        comparison_files = list(output_dir.glob("*comparison_summary*.json"))
        if comparison_files:
            with open(comparison_files[0], "r", encoding="utf-8") as f:
                result["comparison"] = json.load(f)

        s["sim_result"] = result
        s["sim_stage"] = "done"
        cb(datetime.now().strftime("%H:%M:%S"), "✅ 仿真完成")
        _sim_done(session_id, loop)
    except Exception as e:
        _sim_done(session_id, loop, f"仿真执行失败: {e}")


# --- 仿真基线/方案保存 ---

def save_sim_baseline(session_id: str, baseline_power_kwh: float) -> dict:
    """保存 old_station 基准工况年耗电量（单位 kWh）"""
    s = session_store.get(session_id)
    if not s:
        raise KeyError(f"session {session_id!r} 不存在")
    if s.get("sim_mode") != "old_station":
        raise ValueError("仅 old_station 模式允许设置基准耗电量")
    if baseline_power_kwh is None or baseline_power_kwh <= 0:
        raise ValueError("基准耗电量必须 > 0")
    s["sim_baseline_power_kwh"] = float(baseline_power_kwh)
    return s


def save_sim_schemes(session_id: str, schemes: list) -> dict:
    """保存 old_station 优化方案列表"""
    s = session_store.get(session_id)
    if not s:
        raise KeyError(f"session {session_id!r} 不存在")
    if s.get("sim_mode") != "old_station":
        raise ValueError("仅 old_station 模式允许设置优化方案")
    if not isinstance(schemes, list):
        raise ValueError("schemes 必须是数组")
    # 字段校验
    required = {"id", "name", "description", "target_saving_rate_pct"}
    seen_ids = set()
    for idx, sch in enumerate(schemes):
        missing = required - set(sch.keys())
        if missing:
            raise ValueError(f"方案[{idx}] 缺少字段: {missing}")
        if not (0 < sch["target_saving_rate_pct"] < 100):
            raise ValueError(f"方案[{idx}] target_saving_rate_pct 必须在 (0, 100) 之间")
        if sch["id"] in seen_ids:
            raise ValueError(f"方案 id 重复: {sch['id']}")
        seen_ids.add(sch["id"])
    s["sim_schemes"] = schemes
    return s


# --- 启动仿真任务的入口函数 ---

def start_sim_topology_extract(session_id: str, upload_path: Path, main_loop) -> None:
    ensure_ai_queue(session_id)
    session_store._main_loop = main_loop
    _start_sim_task(session_id, _run_topology_extraction, main_loop, upload_path=str(upload_path), loop=main_loop)


def start_sim_topology_parse(session_id: str, text_input: str, project_name: str, main_loop) -> None:
    ensure_ai_queue(session_id)
    session_store._main_loop = main_loop
    _start_sim_task(session_id, _run_topology_parse, main_loop, text_input=text_input, project_name=project_name, loop=main_loop)


def start_sim_cop_parse(session_id: str, img_path: Path, idx: int, model_name: str, rated_cop: float, main_loop) -> None:
    ensure_ai_queue(session_id)
    session_store._main_loop = main_loop
    _start_sim_task(session_id, _run_cop_parse, main_loop, img_path=str(img_path), idx=idx, model_name=model_name, rated_cop=rated_cop, loop=main_loop)


def start_sim_operating_generate(session_id: str, text_input: str, main_loop) -> None:
    ensure_ai_queue(session_id)
    session_store._main_loop = main_loop
    _start_sim_task(session_id, _run_operating_generate, main_loop, text_input=text_input, loop=main_loop)


def start_sim_execute(session_id: str, main_loop) -> None:
    ensure_ai_queue(session_id)
    session_store._main_loop = main_loop
    _start_sim_task(session_id, _run_simulation, main_loop, loop=main_loop)


# ==================== 仿真辅助函数（从 simulation_ui.py 迁移） ====================

_INVALID_TEXT_PATTERNS = [
    "例如：", "示例:", "选填", "选框勾选", "是否", "有没有",
    "请注意", "举例", "需要让运维", "询问运维", "各设备铭牌",
    "描述一下", "客户", "末端", "该项目", "涉及", "整体情况",
    "运行规律", "特殊情况说明", "设备启停方式", "末端设备数量",
    "通讯接口", "供水温度", "运行台数", "运行频率", "出水温度",
    "工频运行则为", "变频运行可填写", "整个冷站", "填写大致的范围",
]


def _is_real_equipment(entry: dict) -> bool:
    """判断一条记录是否是真实设备数据"""
    brand = str(entry.get("brand", "") or "")
    model = str(entry.get("model", "") or "")
    if len(brand) > 40:
        return False
    combined = brand + model
    for pattern in _INVALID_TEXT_PATTERNS:
        if pattern in combined:
            return False
    capacity = entry.get("capacity_kw") or entry.get("额定制冷量") or entry.get("制冷量") or 0
    power = entry.get("power_kw") or entry.get("额定功率") or entry.get("功率") or 0
    cop = entry.get("rated_cop") or entry.get("COP") or entry.get("cop") or 0
    fan_power = entry.get("fan_power_kw") or entry.get("风机功率") or 0
    try:
        if float(capacity) > 0 or float(power) > 0 or float(cop) > 0 or float(fan_power) > 0:
            return True
    except (ValueError, TypeError):
        pass
    if brand and model:
        return True
    if brand and len(brand) <= 20 and not any(p in brand for p in ["类型", "否", "是", "/"]):
        return True
    if model and len(model) <= 60:
        return True
    return False


def _normalize_topology(raw: dict) -> dict:
    """标准化 AI 输出的拓扑 JSON"""
    if "topology" in raw and isinstance(raw["topology"], dict):
        inner = raw["topology"]
        project_name = raw.get("project_name", inner.get("project_name", raw.get("项目名称", "未知项目")))
        raw = inner
    else:
        project_name = raw.get("project_name", raw.get("项目名称", "未知项目"))

    topology = {"project_name": project_name}

    # 冷水主机
    chillers = (
        raw.get("chillers") or raw.get("冷水主机") or raw.get("主机") or
        raw.get("chiller") or raw.get("冷水机组") or []
    )
    normalized_chillers = []
    for ch in chillers:
        if not isinstance(ch, dict) or not _is_real_equipment(ch):
            continue
        normalized_chillers.append({
            "brand": ch.get("brand", ch.get("品牌", "")) or "",
            "model": ch.get("model", ch.get("型号", "")) or "",
            "type": ch.get("type", ch.get("类型", "")) or "",
            "count": int(ch.get("count", ch.get("台数", ch.get("数量", 1))) or 1),
            "capacity_kw": float(ch.get("capacity_kw", ch.get("额定制冷量", ch.get("制冷量", 0))) or 0),
            "power_kw": float(ch.get("power_kw", ch.get("额定功率", ch.get("功率", 0))) or 0),
            "rated_cop": float(ch.get("rated_cop", ch.get("COP", ch.get("额定COP", ch.get("cop", 0)))) or 0),
            "vfd": ch.get("vfd", ch.get("变频", True)),
        })

    # 纠正水泵误分类
    misclassified_pumps = []
    real_chillers = []
    for ch in normalized_chillers:
        cop = ch.get("rated_cop", 0)
        power = ch.get("power_kw", 0)
        capacity = ch.get("capacity_kw", 0)
        computed_ratio = capacity / power if power > 0 else 0
        if cop > 10 or (computed_ratio > 10 and power <= 200):
            misclassified_pumps.append(ch)
        else:
            real_chillers.append(ch)
    topology["chillers"] = real_chillers

    extra_chilled_pumps = [
        {"power_kw": p.get("power_kw", 0), "count": p.get("count", 1), "vfd": p.get("vfd", True)}
        for p in misclassified_pumps
    ]

    # 冷却泵
    cooling_pumps = raw.get("cooling_pumps") or raw.get("冷却泵") or []
    topology["cooling_pumps"] = [
        {
            "brand": p.get("brand", p.get("品牌", "")) or "",
            "model": p.get("model", p.get("型号", "")) or "",
            "power_kw": float(p.get("power_kw", p.get("功率", 0)) or 0),
            "count": int(p.get("count", p.get("台数", 1)) or 1),
            "vfd": p.get("vfd", p.get("变频", True)),
            "flow_rate_m3h": float(p.get("flow_rate_m3h", p.get("流量", p.get("额定流量", 0))) or 0),
            "head_m": float(p.get("head_m", p.get("扬程", p.get("压头", 0))) or 0),
        }
        for p in cooling_pumps if isinstance(p, dict) and _is_real_equipment(p)
    ]

    # 冷冻泵
    chilled_pumps = raw.get("chilled_pumps") or raw.get("冷冻泵") or []
    normalized_chilled = [
        {
            "brand": p.get("brand", p.get("品牌", "")) or "",
            "model": p.get("model", p.get("型号", "")) or "",
            "power_kw": float(p.get("power_kw", p.get("功率", 0)) or 0),
            "count": int(p.get("count", p.get("台数", 1)) or 1),
            "vfd": p.get("vfd", p.get("变频", True)),
            "flow_rate_m3h": float(p.get("flow_rate_m3h", p.get("流量", p.get("额定流量", 0))) or 0),
            "head_m": float(p.get("head_m", p.get("扬程", p.get("压头", 0))) or 0),
        }
        for p in chilled_pumps if isinstance(p, dict) and _is_real_equipment(p)
    ]
    existing_powers = {p["power_kw"] for p in normalized_chilled}
    for ep in extra_chilled_pumps:
        if ep["power_kw"] not in existing_powers:
            normalized_chilled.append(ep)
    topology["chilled_pumps"] = normalized_chilled

    # 冷却塔
    towers = raw.get("towers") or raw.get("冷却塔") or []
    topology["towers"] = [
        {
            "brand": t.get("brand", t.get("品牌", "")) or "",
            "model": t.get("model", t.get("型号", "")) or "",
            "fan_power_kw": float(t.get("fan_power_kw", t.get("风机功率", t.get("功率", 0))) or 0),
            "count": int(t.get("count", t.get("台数", 1)) or 1),
            "vfd": t.get("vfd", t.get("变频", True)),
            "type": t.get("type", t.get("类型", "开式")),
            "flow_rate_m3h": float(t.get("flow_rate_m3h", t.get("流量", t.get("额定流量", 0))) or 0),
        }
        for t in towers if isinstance(t, dict) and _is_real_equipment(t)
    ]

    # 特殊结构
    special = raw.get("special") or raw.get("特殊结构") or {"type": "none"}
    if isinstance(special, dict):
        sp_type = special.get("type", "none")
        # 中文类型 → 英文标准键映射
        cn_type_map = {
            "水蓄冷": "storage", "蓄冷": "storage", "水蓄能": "storage",
            "冰蓄冷": "ice", "蓄冰": "ice", "冰蓄能": "ice",
            "热回收": "heat_recovery", "热回收+蓄冷": "storage_and_heat_recovery",
            "无": "none", "none": "none", "None": "none",
        }
        if isinstance(sp_type, str) and sp_type in cn_type_map:
            sp_type = cn_type_map[sp_type]
        # 清理 type 字段：如果包含垃圾文本，尝试关键词匹配
        if isinstance(sp_type, str) and len(sp_type) > 20:
            detected = []
            if "热回收" in sp_type:
                detected.append("heat_recovery")
            if "蓄冷" in sp_type or "蓄冰" in sp_type:
                detected.append("storage")
            sp_type = "_and_".join(detected) if detected else "none"
        special["type"] = sp_type
    elif isinstance(special, str):
        special = {"type": special if len(special) <= 20 else "none"}
    else:
        special = {"type": "none"}
    topology["special"] = special

    return topology


def _generate_default_cop_table(rated_cop: float) -> dict:
    """基于额定 COP 生成典型部分负荷曲线"""
    cw_temps = list(range(18, 34))
    load_ratios = [1.0, 0.9, 0.8, 0.7, 0.6, 0.5, 0.4, 0.3, 0.2]
    plr_factors = [1.00, 1.08, 1.14, 1.18, 1.15, 1.08, 0.95, 0.78, 0.58]
    base_cw_temp = 30.0

    cop_table = []
    for cw_t in cw_temps:
        row = []
        temp_factor = 1.0 - 0.03 * (cw_t - base_cw_temp)
        for plr_f in plr_factors:
            cop = rated_cop * plr_f * temp_factor
            row.append(round(max(cop, 2.0), 2))
        cop_table.append(row)

    return {
        "cw_temps": cw_temps,
        "load_ratios": load_ratios,
        "cop_table": cop_table,
    }


# ==================== 报价模块服务 ====================

def _run_quote_extract_in_thread(
    session_id: str,
    survey_path: Path,
) -> None:
    """在线程中调用 quote_module.ai_client 完成调研表 AI 解析"""
    loop = session_store._main_loop

    def ai_callback(message: str) -> None:
        timestamp = datetime.now().strftime("%H:%M:%S")
        session_store.append_ai_log(session_id, timestamp, message)
        event = {"type": "log", "timestamp": timestamp, "message": message}
        q = get_ai_queue(session_id)
        if q is not None:
            loop.call_soon_threadsafe(q.put_nowait, event)

    s = session_store.get(session_id)
    if not s:
        return

    try:
        ai_callback("🚀 开始 AI 解析调研表...")

        from quote_module.ai_client import read_survey_with_ai
        from apps.api.quote_local_parser import enrich_project_data

        project_data = read_survey_with_ai(
            str(survey_path),
            callback=ai_callback,
        )
        if not project_data:
            raise ValueError("AI 返回空数据")
        project_data = enrich_project_data(project_data)

        s["quote_project_data"] = project_data
        s["quote_parse_mode"] = "ai"
        session_store.update_stage(session_id, "extracted")
        dw_count = len(project_data.get("data_warnings", []))
        ai_callback("✅ AI 解析完成" + (f"（⚠️ {dw_count} 项数据提醒）" if dw_count else ""))

        event = {"type": "complete", "timestamp": datetime.now().strftime("%H:%M:%S"), "message": "提取完成"}
        q = get_ai_queue(session_id)
        if q is not None:
            loop.call_soon_threadsafe(q.put_nowait, event)

    except Exception as e:
        ai_callback(f"⚠️ AI 解析失败: {e}，尝试本地降级解析...")
        try:
            from apps.api.quote_local_parser import read_survey_local, enrich_project_data
            project_data = read_survey_local(str(survey_path))
            project_data = enrich_project_data(project_data)
            s["quote_project_data"] = project_data
            s["quote_parse_mode"] = "local"
            session_store.update_stage(session_id, "extracted")
            dw_count = len(project_data.get("data_warnings", []))
            ai_callback("✅ 本地降级解析完成" + (f"（⚠️ {dw_count} 项数据提醒）" if dw_count else ""))

            event = {"type": "complete", "timestamp": datetime.now().strftime("%H:%M:%S"), "message": "提取完成（本地降级）"}
            q = get_ai_queue(session_id)
            if q is not None:
                loop.call_soon_threadsafe(q.put_nowait, event)

        except Exception as e2:
            session_store.update_stage(session_id, "error")
            s["quote_error"] = f"解析失败: AI={e}, 本地={e2}"
            ai_callback(f"❌ 解析全部失败: {e2}")

            event = {"type": "error", "timestamp": datetime.now().strftime("%H:%M:%S"), "message": str(e2)}
            q = get_ai_queue(session_id)
            if q is not None:
                loop.call_soon_threadsafe(q.put_nowait, event)


def start_quote_extract(session_id: str, main_loop) -> None:
    """启动报价调研表 AI 解析后台线程"""
    s = session_store.get(session_id)
    if not s:
        raise RuntimeError(f"session {session_id!r} 不存在")

    survey_path = s.get("quote_survey_path")
    if not survey_path:
        raise RuntimeError("未设置调研表路径")

    session_store._main_loop = main_loop
    ensure_ai_queue(session_id)
    s["ai_logs"] = []
    q = get_ai_queue(session_id)
    while q is not None and not q.empty():
        q.get_nowait()
    session_store.update_stage(session_id, "extracting")

    t = threading.Thread(
        target=_run_quote_extract_in_thread,
        args=(session_id, Path(survey_path)),
        daemon=True,
    )
    _quote_extract_tasks[session_id] = t
    t.start()


def _run_quote_calculate_in_thread(session_id: str) -> None:
    """在线程中执行报价计算：section9 分析 + calculate"""
    loop = session_store._main_loop

    def ai_callback(message: str) -> None:
        timestamp = datetime.now().strftime("%H:%M:%S")
        session_store.append_ai_log(session_id, timestamp, message)
        event = {"type": "log", "timestamp": timestamp, "message": message}
        q = get_ai_queue(session_id)
        if q is not None:
            loop.call_soon_threadsafe(q.put_nowait, event)

    s = session_store.get(session_id)
    if not s:
        return

    try:
        ai_callback("🔢 开始计算报价...")

        project_data = s.get("quote_project_data")
        if not project_data:
            raise RuntimeError("缺少项目数据，请先完成提取")

        from quote_module.section9_analyzer import analyze_section9
        from quote_module.scripts.calculate import calculate as quote_calc

        notes = project_data.get("project_notes", {})
        current_status = notes.get("current_status", "")

        s9_parts = []
        for arr_field in ("smart_analysis", "special_requirements", "other_notes"):
            arr = notes.get(arr_field, [])
            if isinstance(arr, list):
                s9_parts.extend([x for x in arr if x])
            elif isinstance(arr, str) and arr:
                s9_parts.append(arr)
        if not s9_parts and notes.get("goals"):
            s9_parts.append(notes["goals"])
        s9_text = "\n".join(s9_parts)

        def log_callback(message: str) -> None:
            ai_callback(f"[section9] {message}")

        ai_callback("📋 分析项目需求 (section9)...")
        section9_result = analyze_section9(current_status, s9_text, callback=log_callback)
        s["quote_section9_result"] = section9_result

        ai_callback("💰 执行价格计算...")
        catalog_path = PROJECT_ROOT / "quote_module" / "outputs" / "price_catalog.json"
        bottom_sheet = quote_calc(
            project_data,
            catalog_path=str(catalog_path),
            tax_mode="exclusive",
            section9_analysis=section9_result,
        )

        s["quote_bottom_sheet"] = bottom_sheet

        bottom_path = s["session_dir"] / f"bottom_sheet_{session_id}.json"
        with open(bottom_path, "w", encoding="utf-8") as f:
            json.dump(bottom_sheet, f, ensure_ascii=False, indent=2)

        session_store.update_stage(session_id, "calculated")
        ai_callback("✅ 计算完成")

        event = {"type": "complete", "timestamp": datetime.now().strftime("%H:%M:%S"), "message": "计算完成"}
        q = get_ai_queue(session_id)
        if q is not None:
            loop.call_soon_threadsafe(q.put_nowait, event)

    except Exception as e:
        session_store.update_stage(session_id, "error")
        s["quote_error"] = str(e)
        ai_callback(f"❌ 计算失败: {e}")
        event = {"type": "error", "timestamp": datetime.now().strftime("%H:%M:%S"), "message": f"计算失败: {e}"}
        q = get_ai_queue(session_id)
        if q is not None:
            loop.call_soon_threadsafe(q.put_nowait, event)


def start_quote_calculate(session_id: str, main_loop) -> None:
    """启动报价计算后台线程"""
    s = session_store.get(session_id)
    if not s:
        raise RuntimeError(f"session {session_id!r} 不存在")

    project_data = s.get("quote_project_data")
    if not project_data:
        raise RuntimeError("缺少项目数据，请先完成提取")

    session_store._main_loop = main_loop
    ensure_ai_queue(session_id)
    s["ai_logs"] = []
    q = get_ai_queue(session_id)
    while q is not None and not q.empty():
        q.get_nowait()
    session_store.update_stage(session_id, "calculating")

    t = threading.Thread(
        target=_run_quote_calculate_in_thread,
        args=(session_id,),
        daemon=True,
    )
    _quote_calculate_tasks[session_id] = t
    t.start()


def _run_quote_generate_in_thread(
    session_id: str,
    bottom_sheet_path: Path,
    output_dir: Path,
) -> None:
    """在线程中生成报价表和方案文档"""
    loop = session_store._main_loop

    def ai_callback(message: str) -> None:
        timestamp = datetime.now().strftime("%H:%M:%S")
        session_store.append_ai_log(session_id, timestamp, message)
        event = {"type": "log", "timestamp": timestamp, "message": message}
        q = get_ai_queue(session_id)
        if q is not None:
            loop.call_soon_threadsafe(q.put_nowait, event)

    s = session_store.get(session_id)
    if not s:
        return

    try:
        ai_callback("🚀 开始生成报价表和方案文档...")

        from quote_module.scripts.generate_quote import generate as gen_quote
        from quote_module.scripts.generate_solution import generate as gen_solution

        output_dir.mkdir(parents=True, exist_ok=True)

        ai_callback("📊 生成报价表 (标准价 + 渠道价)...")
        gen_quote(str(bottom_sheet_path), str(output_dir))

        ai_callback("📝 生成技术方案文档...")
        gen_solution(str(bottom_sheet_path), str(output_dir / "技术方案.docx"))

        generated_files = sorted([str(p) for p in output_dir.glob("*") if p.is_file()])
        s["quote_generated_files"] = generated_files
        session_store.update_stage(session_id, "done")

        ai_callback(f"✅ 生成完成，共 {len(generated_files)} 个文件")

        event = {"type": "complete", "timestamp": datetime.now().strftime("%H:%M:%S"), "message": "生成完成"}
        q = get_ai_queue(session_id)
        if q is not None:
            loop.call_soon_threadsafe(q.put_nowait, event)

    except Exception as e:
        session_store.update_stage(session_id, "error")
        s["quote_error"] = str(e)
        ai_callback(f"❌ 生成失败: {e}")

        event = {"type": "error", "timestamp": datetime.now().strftime("%H:%M:%S"), "message": str(e)}
        q = get_ai_queue(session_id)
        if q is not None:
            loop.call_soon_threadsafe(q.put_nowait, event)


def start_quote_generate(session_id: str, main_loop) -> None:
    """启动报价生成后台线程"""
    s = session_store.get(session_id)
    if not s:
        raise RuntimeError(f"session {session_id!r} 不存在")

    bottom_sheet = s.get("quote_bottom_sheet")
    if not bottom_sheet:
        raise RuntimeError("缺少计算结果，请先完成计算")

    session_store._main_loop = main_loop
    ensure_ai_queue(session_id)
    s["ai_logs"] = []
    q = get_ai_queue(session_id)
    while q is not None and not q.empty():
        q.get_nowait()
    session_store.update_stage(session_id, "generating")

    bottom_sheet_path = s["session_dir"] / f"bottom_sheet_{session_id}.json"
    output_dir = s["session_dir"] / f"generated_{session_id}"

    t = threading.Thread(
        target=_run_quote_generate_in_thread,
        args=(session_id, bottom_sheet_path, output_dir),
        daemon=True,
    )
    _quote_generate_tasks[session_id] = t
    t.start()

