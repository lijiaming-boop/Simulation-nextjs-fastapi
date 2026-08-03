"""AI 月度运行参数推荐 — Streamlit UI 组件

自包含的 Streamlit UI 组件，可嵌入仿真输入页面。
调用 render_ai_monthly_recommendation() 即可渲染完整推荐流程。
"""

import io
import json
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
import pandas as pd
import streamlit as st

from .presets import PRESET_SCHEDULES, SCHEDULE_LABELS
from .calendar_parser import (
    DEFAULT_DAY_TYPE,
    parse_calendar_smart,
)
from .recommend import (
    ai_stage1_from_text,
    ai_stage2_recommend,
    ai_parse_calendar,
)

# ═══════════════════════════════════════════════════════════════
#  常量
# ═══════════════════════════════════════════════════════════════

REGIONS = ["华南", "华东", "华中", "华北", "西南", "西北", "东北", "未知"]
MONTH_NAMES = ["1月", "2月", "3月", "4月", "5月", "6月", "7月", "8月", "9月", "10月", "11月", "12月"]
HOUR_LABELS = [f"{h}:00" for h in range(24)]


# ═══════════════════════════════════════════════════════════════
#  状态初始化
# ═══════════════════════════════════════════════════════════════

def _init_state():
    """初始化本模块的 session_state。"""
    defaults = {
        "aim_stage_one_done": False,
        "aim_survey_text": None,
        "aim_analysis": None,
        "aim_schedule_label": "单班制办公",
        "aim_workday_curve": None,
        "aim_offday_curve": None,
        "aim_user_region": None,
        "aim_total_chiller_count": 3,
        "aim_day_type": DEFAULT_DAY_TYPE,
        "aim_day_type_rule": "",
        "aim_stage_two_done": False,
        "aim_monthly_chiller_count": [1] * 12,
        "aim_monthly_load_factors": [0.5] * 12,
        "aim_summary": "",
    }
    for k, v in defaults.items():
        if k not in st.session_state:
            st.session_state[k] = v


def _apply_preset_curves(label: str):
    """将指定班制的预设曲线写入 session_state。"""
    preset = PRESET_SCHEDULES.get(label, PRESET_SCHEDULES["单班制办公"])
    st.session_state.aim_workday_curve = list(preset["workday_curve"])
    st.session_state.aim_offday_curve = list(preset["offday_curve"])
    st.session_state.aim_schedule_label = label


def _extract_text_from_excel(file_bytes: bytes) -> str:
    """从 Excel 提取文本内容。"""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        parts.append(f"\n=== Sheet: {sheet_name} ===\n")
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row,
                                max_col=min(ws.max_column, 20), values_only=True):
            row_text = " | ".join(str(c).strip() if c is not None else "" for c in row)
            if row_text.replace(" | ", "").strip():
                parts.append(row_text)
    wb.close()
    return "\n".join(parts)


# ═══════════════════════════════════════════════════════════════
#  曲线编辑器
# ═══════════════════════════════════════════════════════════════

def _editable_curve(label: str, values: List[float], key: str) -> List[float]:
    """可编辑曲线组件：Plotly 图表 + 可折叠编辑区。"""
    import plotly.graph_objects as go

    curve_key = f"aim_curve_{key}"
    if curve_key not in st.session_state:
        st.session_state[curve_key] = [round(float(v), 2) for v in values]

    current = st.session_state[curve_key]

    fig = go.Figure()
    fig.add_trace(go.Scatter(
        x=list(range(24)), y=current, mode="lines+markers",
        line=dict(width=3, color="#4A90D9"),
        marker=dict(size=10, color="#4A90D9", line=dict(width=1.5, color="white")),
        hovertemplate="<b>%{x}:00</b><br>系数 = %{y:.2f}<extra></extra>",
    ))
    fig.update_layout(
        height=280, margin=dict(l=30, r=10, t=20, b=35),
        xaxis=dict(title="小时", dtick=2, range=[-0.5, 23.5], gridcolor="#f0f0f0"),
        yaxis=dict(title="负载系数", range=[-0.05, 1.25], dtick=0.1, gridcolor="#f0f0f0"),
        title=label, dragmode=False,
    )
    st.plotly_chart(fig, use_container_width=True, key=f"aim_plot_{key}")

    with st.expander("编辑逐时系数", expanded=False):
        changed = False
        for row_idx in range(6):
            cols = st.columns(4)
            for j in range(4):
                i = row_idx * 4 + j
                if i >= 24:
                    break
                with cols[j]:
                    v = st.number_input(
                        f"{i}:00", min_value=0.0, max_value=1.20,
                        value=current[i], step=0.01, format="%.2f",
                        key=f"aim_curve_{key}_{i}", label_visibility="visible",
                    )
                    if abs(v - current[i]) > 0.001:
                        st.session_state[curve_key][i] = round(float(v), 2)
                        changed = True
        if changed:
            st.rerun()

    return st.session_state[curve_key]


# ═══════════════════════════════════════════════════════════════
#  主入口
# ═══════════════════════════════════════════════════════════════

def render_ai_monthly_recommendation() -> Optional[Dict[str, Any]]:
    """渲染 AI 月度运行参数推荐完整流程。

    可在 Streamlit 任何位置调用（expander 内或独立页面均可）。

    Returns:
        推荐结果 dict 供外部读取，包含：
        - monthly_chiller_count, monthly_load_factors
        - workday_curve, offday_curve, day_type
        - schedule_label, user_region, analysis
    """
    _init_state()

    # ── 上传调研表 ──
    st.markdown("### 📎 上传调研表（可选）")
    st.caption("若不提供，AI 将以默认参数分析，后续可手动调整。")
    uploaded = st.file_uploader("上传调研表（Excel）", type=["xlsx", "xls"], key="aim_upload")
    if uploaded:
        with st.spinner("读取中..."):
            try:
                text = _extract_text_from_excel(uploaded.getvalue())
                st.success(f"读取成功，{len(text)} 字符")
                st.session_state.aim_survey_text = text
            except Exception as e:
                st.error(f"读取失败: {e}")

    # ── 阶段一按钮 ──
    st.divider()
    if st.button("🔍 阶段一：AI 分析项目特征", type="primary", use_container_width=True, key="aim_stage1_btn"):
        survey_text = st.session_state.get("aim_survey_text")
        with st.spinner("AI 正在分析..."):
            if survey_text:
                try:
                    result = ai_stage1_from_text(survey_text)
                except Exception as e:
                    st.error(f"AI 调用失败: {e}")
                    result = None
            else:
                result = {
                    "analysis": {
                        "project_type": "无法识别",
                        "project_type_reason": "未上传调研表",
                        "region": "未知",
                        "schedule_label": "单班制办公",
                        "schedule_reason": "未上传调研表，默认班制",
                        "total_chiller_count": 3,
                        "can_identify": False,
                        "summary": "未提供调研表，请手动选择地区和班制。",
                    }
                }

            if result:
                analysis = result.get("analysis", {})
                st.session_state.aim_analysis = analysis
                st.session_state.aim_user_region = analysis.get("region", "未知")
                st.session_state.aim_total_chiller_count = analysis.get("total_chiller_count", 3)
                label = analysis.get("schedule_label", "单班制办公")
                if label not in SCHEDULE_LABELS:
                    label = "单班制办公"
                _apply_preset_curves(label)
                st.session_state.aim_stage_one_done = True
                st.success("✅ 阶段一完成！请核查下方地区和曲线。")
            else:
                st.error("AI 解析失败，请重试")

    # ── 阶段一结果区域 ──
    if not st.session_state.get("aim_stage_one_done") or not st.session_state.get("aim_analysis"):
        return None

    analysis = st.session_state.aim_analysis

    # 分析摘要
    st.divider()
    st.markdown("### 🔍 AI 分析结果")
    st.markdown(f"""
    | 特征 | AI 判断 |
    |------|--------|
    | 项目类型 | **{analysis.get('project_type', '—')}** |
    | 项目地区 | {analysis.get('region', '—')} |
    | 推荐班制 | **{analysis.get('schedule_label', '—')}** |
    | 主机台数 | {st.session_state.aim_total_chiller_count} 台 |
    """)
    if not analysis.get("can_identify", True):
        st.warning("⚠️ AI 无法识别具体环境，请在下方手动选择地区和班制。")

    st.divider()

    # ── 地区选择 + 班制 + 曲线 ──
    st.markdown("### ⏱️ 项目地区 & 运行班制 & 逐时负载曲线")
    st.caption("地区以用户选择为准，阶段二将以此地区推荐月度参数。")

    # 地区
    ai_region = analysis.get("region", "未知")
    cur_region = st.session_state.get("aim_user_region") or ai_region
    if cur_region not in REGIONS:
        cur_region = "未知"
    new_region = st.selectbox("项目地区", REGIONS, index=REGIONS.index(cur_region), key="aim_region")
    if new_region != st.session_state.get("aim_user_region"):
        st.session_state.aim_user_region = new_region
        st.session_state.aim_stage_two_done = False  # 地区变了，需重新推荐

    st.divider()

    # 班制
    cur_label = st.session_state.get("aim_schedule_label", "单班制办公")
    if cur_label not in SCHEDULE_LABELS:
        cur_label = "单班制办公"
    prev_label = st.session_state.get("_aim_prev_label", "")
    new_label = st.selectbox("运行班制", SCHEDULE_LABELS, index=SCHEDULE_LABELS.index(cur_label),
                              key="aim_schedule_select")
    if new_label != prev_label:
        _apply_preset_curves(new_label)
        for k in list(st.session_state.keys()):
            if k.startswith("aim_curve_"):
                del st.session_state[k]
        st.session_state._aim_prev_label = new_label
        st.session_state.aim_stage_two_done = False
        st.rerun()

    preset = PRESET_SCHEDULES[new_label]
    st.caption(f"📝 {preset['desc']}")

    st.divider()

    # 工作日曲线
    st.markdown("#### 📈 工作日逐时负载系数")
    wd = _editable_curve("工作日", st.session_state.aim_workday_curve, f"workday_{new_label}")
    st.session_state.aim_workday_curve = wd

    st.divider()

    # 休息日曲线
    st.markdown("#### 📉 休息日逐时负载系数")
    od = _editable_curve("休息日", st.session_state.aim_offday_curve, f"offday_{new_label}")
    st.session_state.aim_offday_curve = od

    # ── 工作日历规则 ──
    st.divider()
    st.markdown("#### 📅 工作日历规则（可选）")
    col_t, col_b = st.columns([5, 1])
    with col_t:
        cal_text = st.text_input(
            "运行日历规则",
            placeholder="默认五天两休，支持修改工作周期  例：上六休一，考虑所有节假日（法定），端午上班",
            key="aim_cal_rule", label_visibility="collapsed",
        )
    with col_b:
        parse_btn = st.button("📅 解析", key="aim_cal_parse", use_container_width=True)

    if parse_btn and cal_text.strip():
        result = parse_calendar_smart(cal_text.strip())
        if result is None or not result.get("day_type"):
            # 程序化失败，回退 AI
            st.info("程序化解析未匹配，尝试 AI 解析...")
            ai_dt = ai_parse_calendar(cal_text.strip())
            if ai_dt:
                result = {"day_type": ai_dt, "method": "ai", "pattern": None, "holidays": []}
        if result and result.get("day_type"):
            day_type = result["day_type"]
            st.session_state.aim_day_type = day_type
            st.session_state.aim_day_type_rule = cal_text.strip()
            work_n = day_type.count(0)
            rest_n = day_type.count(1)
            detail = []
            if result.get("pattern"):
                detail.append(f"循环：{result['pattern']}")
            if result.get("holidays"):
                detail.append(f"节假日休息：{'、'.join(result['holidays'])}")
            if result.get("excluded_holidays"):
                detail.append(f"排除：{'、'.join(result['excluded_holidays'])}")
            if result.get("work_on_holidays"):
                detail.append(f"上班：{'、'.join(result['work_on_holidays'])}")
            extra = "  |  " + " · ".join(detail) if detail else ""
            st.success(f"✅ 日历已生成：{work_n} 工作日 + {rest_n} 休息日（后续将依此决定每日负载曲线）{extra}")
        else:
            st.warning("无法解析，请换一种表达方式。")
    elif parse_btn and not cal_text.strip():
        st.warning("请先输入日历规则")

    # ── 阶段二 ──
    st.divider()
    st.markdown("### 🧠 阶段二：AI 推荐月度参数")

    if st.button("📊 确认曲线，推荐月度参数", type="primary", use_container_width=True, key="aim_stage2_btn"):
        tc = st.session_state.aim_total_chiller_count
        ur = st.session_state.get("aim_user_region")
        with st.spinner("AI 正在推荐月度参数..."):
            try:
                r2 = ai_stage2_recommend(
                    analysis, new_label, wd, od, tc, user_region=ur,
                )
            except Exception as e:
                st.error(f"AI 调用失败: {e}")
                r2 = None

        if r2:
            st.session_state.aim_monthly_chiller_count = r2.get("monthly_chiller_count", [1] * 12)
            st.session_state.aim_monthly_load_factors = r2.get("monthly_load_factors", [0.5] * 12)
            st.session_state.aim_summary = r2.get("recommendation_summary", "")
            st.session_state.aim_stage_two_done = True
            st.success("✅ 阶段二完成！下方可手动微调月度参数。")
        else:
            st.error("AI 解析失败，请重试")

    # ── 月度参数编辑器 ──
    if st.session_state.get("aim_stage_two_done"):
        st.divider()

        mc = list(st.session_state.aim_monthly_chiller_count) + [1] * 12
        mf = list(st.session_state.aim_monthly_load_factors) + [0.5] * 12

        st.markdown("### 📋 月度运行参数（可手动修改）")

        st.caption("每月开机台数")
        ec = [0] * 12
        cols_mc = st.columns(12)
        for i, col in enumerate(cols_mc):
            with col:
                ec[i] = col.number_input(MONTH_NAMES[i], 0, 10, min(mc[i], 10), key=f"aim_mc_{i}")

        st.caption("月度负载系数")
        ef = [0.0] * 12
        cols_mf = st.columns(12)
        for i, col in enumerate(cols_mf):
            with col:
                ef[i] = col.number_input(MONTH_NAMES[i], 0.0, 1.0, round(float(mf[i]), 2),
                                         0.01, format="%.2f", key=f"aim_mf_{i}")

        st.session_state.aim_monthly_chiller_count = ec
        st.session_state.aim_monthly_load_factors = ef

        if st.session_state.get("aim_summary"):
            st.info(f"💡 AI 推荐理由：{st.session_state.aim_summary}")

        # ── 导出 ──
        st.divider()
        st.markdown("### 📦 导出结果")
        export = _build_export()
        st.json(export)
        st.download_button(
            "⬇️ 下载运行参数 JSON",
            data=json.dumps(export, ensure_ascii=False, indent=2),
            file_name="ai_recommended_operation_params.json",
            mime="application/json", use_container_width=True, key="aim_download",
        )

    return _build_export()


def _build_export() -> Dict[str, Any]:
    """构建导出字典。"""
    return {
        "analysis": st.session_state.get("aim_analysis", {}),
        "user_region": st.session_state.get("aim_user_region", "未知"),
        "schedule": {
            "label": st.session_state.get("aim_schedule_label", "单班制办公"),
            "workday_curve": [round(v, 2) for v in (st.session_state.get("aim_workday_curve") or [])],
            "offday_curve": [round(v, 2) for v in (st.session_state.get("aim_offday_curve") or [])],
            "day_type": st.session_state.get("aim_day_type", DEFAULT_DAY_TYPE),
        },
        "monthly_chiller_count": st.session_state.get("aim_monthly_chiller_count", [1] * 12),
        "monthly_load_factors": [round(f, 2) for f in (st.session_state.get("aim_monthly_load_factors") or [0.5] * 12)],
    }


def get_recommended_params() -> Dict[str, Any]:
    """供外部调用：获取当前推荐参数（不渲染UI）。"""
    return _build_export()
