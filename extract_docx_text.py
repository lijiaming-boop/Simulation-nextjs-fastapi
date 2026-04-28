#!/usr/bin/env python3
"""
提取Word文档内容，查看占位符格式
"""

from docx import Document
from pathlib import Path

def extract_docx_text(docx_path):
    """提取Word文档的文本内容"""
    try:
        doc = Document(docx_path)
        text = []
        
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text.append(paragraph.text)
        
        # 提取表格内容
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    cell_text = ' '.join([p.text for p in cell.paragraphs if p.text.strip()])
                    if cell_text:
                        row_text.append(cell_text)
                if row_text:
                    text.append(' | '.join(row_text))
        
        return '\n'.join(text)
        
    except Exception as e:
        print(f"提取文档内容失败: {str(e)}")
        return ""


def main():
    """主函数"""
    docx_path = Path("00_sample/模板测算表/空压系统数智化升级方案模板.docx")
    
    if not docx_path.exists():
        print(f"文件不存在: {docx_path}")
        return
    
    print(f"提取文档: {docx_path}")
    print("=" * 80)
    
    text = extract_docx_text(docx_path)
    print(text)
    
    print("=" * 80)
    print("提取完成")


if __name__ == "__main__":
    main()