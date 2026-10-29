"""
标准能管方案生成器
从内部底稿 JSON 读取项目信息，复制方案模板并替换内容，删除禁入内容
"""
import json
import re
import sys
from pathlib import Path
from docx import Document


# ===================== 路径配置 =====================

TEMPLATE_PATH = Path(__file__).parent.parent / 'templates' / '标准能管方案模板.docx'

# ===================== 禁入内容规则 =====================

# 整段删除（匹配任一即清空整个段落/单元格）
DELETE_IF_CONTAINS = [
    '高级分析',          # 高级分析（二期可选）
    '二期可选', '二期实现', '二期功能',
    '一期功能需要', '二期在一期', '一期、二期',
    '能源AI预测', 'AI预测', 'AI自动分析',
    'AI可视化', '人工智能', '大模型',
    '自动启停', '自动优化', '闭环控制',
    '固定节能率', '诊断报告', '优化建议',
    '能源计划与趋势分析',
    '分阶段实施',
    '单位产品能耗（二期',
    '能源分析（二期',
]

# 文本级替换（在段落/单元格内部的文本替换，只替换匹配部分）
TEXT_REPLACEMENTS = [
    ('能耗预测分析', '能耗趋势分析'),
    ('能源预测和分析', '能源趋势分析'),
    ('能源预测与分析', '能源趋势分析'),
    ('能源预测', '能源趋势分析'),
    ('预测分析', '趋势分析'),
    ('智能分析', '能耗分析'),
    ('AI', ''),
    ('（二期可选）', ''),
    ('（二期实现）', ''),
    ('（一期实现）', ''),
    ('一期、二期', ''),
    ('二期', ''),
    ('一期', ''),
    ('分阶段实施', '系统性实施'),
    ('自动抄表', '自动抄表'),   # 允许保留
]


def get_project_description(data: dict) -> dict:
    """从 JSON 提取项目描述文本"""
    info = data.get('project_info', {})
    notes = data.get('project_notes', {})
    calcs = data.get('calculations', {})
    opts = data.get('optional_modules', {})
    third_party = data.get('third_party_systems', [])

    client = info.get('client_name', '某企业')
    project = info.get('project_name', '能源管理系统')
    location = info.get('project_location', '')
    deploy = info.get('deployment_mode', 'cloud')

    deploy_text = {'cloud': '云端部署'}.get(deploy, '本地私有化部署')

    # 项目背景
    background = (
        f'{client}计划建设{project}，旨在搭建能源管理协同平台，'
        f'通过信息化手段贯通从采购、输送、分配到使用的全周期能源管理数据链，'
        f'构建能源成本中心，支持自动统计分析、在线监控预警及用能异常的改进协同，'
        f'最终推动能源管理可视化，达到能源管理统一规范、提高过程管控能力、降低成本的目的。\n\n'
        f'{notes.get("current_status", "")}'
    )

    # 项目范围
    total_electric = calcs.get('total_electric_meters', 0)
    total_fluid = calcs.get('total_fluid_meters', 0)
    third_count = calcs.get('third_party_count', 0)

    scope_parts = [f'本项目覆盖{client}的生产车间、辅助设施及办公场所的能源管理。']
    scope_parts.append(f'接入电表共{total_electric}台，水表及流体类表计共{total_fluid}台。')

    if third_party:
        sys_names = '、'.join(s['name'] for s in third_party)
        scope_parts.append(f'对接{third_count}个第三方系统（{sys_names}），实现数据单向读取与集中展示。')

    selected_module_names = [m.get("name", "") for m in opts.get("selected_modules", [])]
    if any("批次" in n for n in selected_module_names):
        scope_parts.append('配置批次单耗分析模块，支持按批次/产品类型分析单位能耗。')
    if any("移动" in n or "APP" in n or "小程序" in n for n in selected_module_names):
        scope_parts.append('配置移动端（微信小程序），支持实时数据查看与预警消息推送。')
    if opts.get('extra_reports', 0) > 0:
        scope_parts.append(f'在标准3张报表基础上，额外定制{opts["extra_reports"]}张报表。')

    scope_parts.append(f'系统采用{deploy_text}方式部署。')
    scope_parts.append('本项目配置1张综合能源大屏、3张组态/能流拓扑页面、3张标准能源报表。')

    scope = '\n'.join(scope_parts)

    # 项目目标（无分期语言）
    goals = (
        f'1. 实现能源管控一体化，达到管理节能的目的\n'
        f'本方案采用系统化采集技术，为{client}提供集计量、监测和管理于一体的能源数字化管理平台，'
        f'从管理节能方面助力能源管理，提升能源管理水平。\n\n'
        f'2. 通过实时在线监测，提高安全管理水平\n'
        f'能源数字化管理平台可对电压、电流等实时信息进行超限报警，也可以对异常能耗进行预警，'
        f'通过告警推送帮助管理人员及时发现异常并跟进处理。\n\n'
        f'3. 通过数字化工具，为能源管理体系提供支撑\n'
        f'能源数字化管理平台数据将为专业管理人员提供科学用能工具，'
        f'为实现总体用能平衡、贯彻能效指标、能耗公示、定额考核等管理节能手段提供支持。'
    )

    # 建设范围
    total_elec = calcs.get('total_electric_meters', 0)
    new_elec_count = sum(
        m.get('quantity', 0) for m in data.get('new_meters', {}).get('electric', [])
    ) if 'new_meters' in data else 0
    existing_elec = total_elec - new_elec_count

    construction_scope = (
        f'经现场调研，本次项目范围为：\n'
        f'· 接入既有可通讯电表约{existing_elec}台（具体以调研清单为准）；\n'
        f'· 新增三相多功能电表{new_elec_count}台及配套互感器；\n'
        f'· 新增水表及配套安装；\n'
        f'· 部署能管云盒、通讯箱等边缘采集设备；\n'
        f'· 搭建能源管理平台软件（含能源大屏、组态/能流拓扑、能源报表）；\n'
        f'· 完成系统调试与用户培训。'
    )

    return {
        'background': background,
        'scope': scope,
        'goals': goals,
        'construction_scope': construction_scope,
        'client': client,
        'project': project,
        'location': location,
    }


def _should_delete(text: str) -> bool:
    """判断文本是否应整段删除"""
    for kw in DELETE_IF_CONTAINS:
        if kw in text:
            return True
    return False


def _clean_text(text: str) -> str:
    """对文本执行替换清理"""
    result = text
    for old, new in TEXT_REPLACEMENTS:
        result = result.replace(old, new)
    return result


def _process_paragraph_runs(para, proj: dict):
    """
    对单个段落的 runs 进行替换处理。
    先全文判断是否需删除；再逐 run 替换占位符和禁词。
    """
    full_text = para.text

    # --- 特殊段落：精确匹配后整体替换 ---
    if '[客户名称]' in full_text and ('计划建设' in full_text or '能源管理系统' in full_text):
        _replace_para(para, proj['background'])
        return

    if '能源管理范围覆盖' in full_text and '生产车间' in full_text:
        _replace_para(para, proj['scope'])
        return

    if '经现场实地调研' in full_text and '全量表计清单' in full_text:
        _replace_para(para, proj['construction_scope'])
        return

    if '分阶段实施' in full_text:
        _replace_para(para, proj['goals'])
        return

    # 一期/二期 描述性段落
    if ('一期功能需要满足' in full_text or '二期功能在一期' in full_text
            or '能源分析（二期' in full_text or '单位产品能耗（二期' in full_text
            or '能源计划与趋势分析（二期' in full_text):
        _clear_para(para)
        return

    # --- 一般处理 ---
    # 检查是否需整段删除
    if _should_delete(full_text):
        _clear_para(para)
        return

    # 逐 run 处理
    for run in para.runs:
        # 替换客户占位符（注意模板中可能有 "[客户名称]有限公司" 的写法）
        run.text = run.text.replace('[客户名称]有限公司', proj['client'])
        run.text = run.text.replace('[客户名称]', proj['client'])
        run.text = run.text.replace('能源数字化管理系统', proj['project'])
        # 替换禁词
        run.text = _clean_text(run.text)


def _process_all_paragraphs(doc, proj: dict):
    """处理文档所有段落"""
    for para in doc.paragraphs:
        _process_paragraph_runs(para, proj)


def _process_all_tables(doc, proj: dict):
    """处理文档所有表格"""
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text = cell.text

                # 整格删除判断
                if _should_delete(full_text):
                    for para in cell.paragraphs:
                        _clear_para(para)
                    continue

                # 逐段落清理
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.text = run.text.replace('[客户名称]有限公司', proj['client'])
                        run.text = run.text.replace('[客户名称]', proj['client'])
                        run.text = run.text.replace('能源数字化管理系统', proj['project'])
                        run.text = _clean_text(run.text)


def generate(bottom_sheet_path: str, output_path: str):
    data = json.loads(Path(bottom_sheet_path).read_text(encoding='utf-8'))
    proj = get_project_description(data)

    # 打开模板
    doc = Document(str(TEMPLATE_PATH))

    # 第一遍：处理段落
    _process_all_paragraphs(doc, proj)

    # 第二遍：处理表格
    _process_all_tables(doc, proj)

    # 第三遍（安全网）：再次扫描所有文本，对仍残留的禁词做最终替换
    _final_cleanup(doc)

    # ---- 保存 ----
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return proj


def _final_cleanup(doc):
    """
    最终安全网：扫描所有段落和表格，对仍存在的禁词做替换。
    不整段删除 —— 只做文本替换，避免误伤。
    """
    # 最终残留替换表（更激进）
    final_replacements = [
        ('（二期）', ''),
        ('二期', ''),
        ('一期', ''),
        ('AI', ''),
        ('人工智能', ''),
        ('大模型', ''),
        ('预测分析', '趋势分析'),
        ('能耗预测', '能耗趋势'),
        ('能源预测', '能源趋势'),
        ('预测', '趋势'),
        ('分阶段', ''),
        ('自动启停', ''),
        ('闭环控制', ''),
        ('固定节能率', ''),
        ('诊断报告', '数据报表'),
        ('优化建议', ''),
    ]

    def clean_runs(paragraph):
        for run in paragraph.runs:
            for old, new in final_replacements:
                run.text = run.text.replace(old, new)

    for para in doc.paragraphs:
        clean_runs(para)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    clean_runs(para)

    # 清理 sections 中的页眉页脚
    for section in doc.sections:
        for para in section.header.paragraphs:
            clean_runs(para)
        for para in section.footer.paragraphs:
            clean_runs(para)


def _replace_para(para, new_text: str):
    """替换段落全部文本 — 清除所有run后写入第一个run"""
    # 先清空所有 run
    for run in para.runs:
        run.text = ''
    # 再写入第一个 run（如果没有 run 则创建）
    if para.runs:
        para.runs[0].text = new_text
    else:
        para.add_run(new_text)


def _clear_para(para):
    """清空段落"""
    if para.runs:
        for run in para.runs:
            run.text = ''
    else:
        para.text = ''


if __name__ == '__main__':
    bottom_path = sys.argv[1] if len(sys.argv) > 1 else 'outputs/internal_bottom_sheet.json'
    out_path = sys.argv[2] if len(sys.argv) > 2 else 'outputs/标准技术方案.docx'

    proj_info = generate(bottom_path, out_path)
    print(f'[OK] 标准技术方案: {out_path}')
    print(f'    客户: {proj_info["client"]}')
    print(f'    项目: {proj_info["project"]}')
